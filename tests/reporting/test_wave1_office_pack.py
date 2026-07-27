"""Regression tests for the run-bound Wave 1 Office reporting pack."""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from src.reporting.build_wave1_office_pack import (
    OfficePackError,
    generate_office_pack,
)


ROOT = Path(__file__).resolve().parents[2]
VALIDATION = (
    ROOT / "outputs/qa/validated/latest/wave1_independent_validation.json"
)
CLAIM_MANIFEST = (
    ROOT / "outputs/final/wave1_claim_manifest__W1-20260727-001.json"
)


def test_real_validated_inputs_generate_reopen_and_cover_all_rows(
    tmp_path: Path,
) -> None:
    result = generate_office_pack(
        VALIDATION,
        CLAIM_MANIFEST,
        tmp_path,
        artifact_root=ROOT,
    )

    assert result.excel_path.exists()
    assert result.powerpoint_path.exists()
    assert result.word_path.exists()
    assert result.manifest_path.exists()
    assert result.manifest["run_id"] == "W1-20260727-001"
    assert result.manifest["publication"]["status"] == "pending"
    values = result.manifest["cross_artifact_values"]
    assert values["signal_bucket_count"] == 23
    assert values["risk_band_count"] == 3
    assert len(values["signal_rows"]) == 23
    assert len(values["risk_band_rows"]) == 3

    workbook = load_workbook(result.excel_path, read_only=False)
    assert workbook["Signal_Buckets"].max_row == 24
    assert workbook["Risk_Bands"].max_row == 4
    assert len(workbook["Risk_Bands"]._charts) == 1
    workbook.close()


def test_second_generation_refuses_to_overwrite(tmp_path: Path) -> None:
    generate_office_pack(
        VALIDATION,
        CLAIM_MANIFEST,
        tmp_path,
        artifact_root=ROOT,
    )

    with pytest.raises(OfficePackError, match="refusing to overwrite"):
        generate_office_pack(
            VALIDATION,
            CLAIM_MANIFEST,
            tmp_path,
            artifact_root=ROOT,
        )


def test_tampered_validation_is_blocked_by_claim_hash(
    tmp_path: Path,
) -> None:
    validation = json.loads(VALIDATION.read_text(encoding="utf-8"))
    validation["verified_findings"]["source_rows"] += 1
    tampered = tmp_path / "tampered-validation.json"
    tampered.write_text(json.dumps(validation), encoding="utf-8")

    with pytest.raises((OfficePackError, ValueError)):
        generate_office_pack(
            tampered,
            CLAIM_MANIFEST,
            tmp_path / "output",
            artifact_root=ROOT,
        )


def test_mixed_run_claim_manifest_is_blocked(tmp_path: Path) -> None:
    manifest = json.loads(CLAIM_MANIFEST.read_text(encoding="utf-8"))
    manifest["claims"][0]["run_id"] = "W1-OTHER-RUN"
    mixed = tmp_path / "mixed-claim-manifest.json"
    mixed.write_text(json.dumps(manifest), encoding="utf-8")

    with pytest.raises(OfficePackError, match="run ID"):
        generate_office_pack(
            VALIDATION,
            mixed,
            tmp_path / "output",
            artifact_root=ROOT,
        )


def test_ooxml_hashes_are_deterministic_across_output_directories(
    tmp_path: Path,
) -> None:
    first = generate_office_pack(
        VALIDATION,
        CLAIM_MANIFEST,
        tmp_path / "first",
        artifact_root=ROOT,
    )
    second = generate_office_pack(
        VALIDATION,
        CLAIM_MANIFEST,
        tmp_path / "second",
        artifact_root=ROOT,
    )

    first_hashes = {
        row["artifact_type"]: row["sha256"]
        for row in first.manifest["artifacts"]
    }
    second_hashes = {
        row["artifact_type"]: row["sha256"]
        for row in second.manifest["artifacts"]
    }
    assert first_hashes == second_hashes
