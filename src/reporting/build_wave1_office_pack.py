#!/usr/bin/env python3
"""Generate and verify the run-bound Wave 1 Excel, PowerPoint, and Word pack.

The independent validation JSON is the numerical authority. The existing claim
manifest is the claim/run/hash authority. Missing, failed, mixed-run, tampered,
or inconsistent evidence blocks every Office output. Generated files remain
internal-only pending Gate 5 review and explicit human publication approval.
"""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import re
import shutil
import tempfile
import zipfile
from dataclasses import dataclass
from datetime import datetime, timezone
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from typing import Any, Mapping, Sequence

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches as DocxInches
from docx.shared import Pt as DocxPt
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from src.reporting.build_wave1_report import build_report_bundle
from src.validation.gate_evidence import load_json_object


REPOSITORY_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_VALIDATION = (
    REPOSITORY_ROOT
    / "outputs/qa/validated/latest/wave1_independent_validation.json"
)
DEFAULT_CLAIM_MANIFEST = (
    REPOSITORY_ROOT / "outputs/final/wave1_claim_manifest__W1-20260727-001.json"
)
DEFAULT_OUTPUT_DIR = REPOSITORY_ROOT / "outputs/final"

INTERNAL_STATUS = (
    "INTERNAL ONLY — PENDING GATE 5 AND HUMAN PUBLICATION APPROVAL"
)
IFRS9_LIMITATION = (
    "Public-data descriptive prototype; not a production IFRS 9 rating, "
    "Stage, PD, LGD, EAD, or ECL result."
)
REQUIRED_CLAIM_IDS = frozenset(
    {
        "W1-CLM-001",
        "W1-CLM-002",
        "W1-CLM-003",
        "W1-CLM-004",
        "W1-CLM-005",
        "W1-CLM-006",
    }
)
FIXED_ZIP_TIMESTAMP = (1980, 1, 1, 0, 0, 0)

NAVY = "17365D"
BLUE = "2E75B6"
LIGHT_BLUE = "D9EAF7"
PALE = "F3F6F9"
AMBER = "F4B183"
GREEN = "70AD47"
DARK = "243447"
WHITE = "FFFFFF"
RED = "C00000"


class OfficePackError(ValueError):
    """Raised when inputs or generated artifacts cannot pass the pack contract."""


@dataclass(frozen=True)
class OfficeData:
    """Validated values shared by all three Office artifacts."""

    run_id: str
    generated_at: str
    generated_datetime: datetime
    definition_version: str
    source_id: str
    source_sha256: str
    validation_path: str
    validation_sha256: str
    claim_manifest_path: str
    claim_manifest_sha256: str
    source_rows: int
    raw_rows: int
    borrower_rows: int
    account_month_rows: int
    months_per_borrower: int
    bill_source_total: str
    bill_derived_total: str
    bill_difference: str
    payment_source_total: str
    payment_derived_total: str
    payment_difference: str
    sql_test_count: int
    sql_failed_count: int
    signal_count: int
    signal_rows: tuple[Mapping[str, Any], ...]
    band_rows: tuple[Mapping[str, Any], ...]
    gate_rows: tuple[Mapping[str, Any], ...]
    dq_rows: tuple[Mapping[str, Any], ...]
    claims: tuple[Mapping[str, Any], ...]
    limitations: tuple[str, ...]


@dataclass(frozen=True)
class OfficePackResult:
    """Paths and manifest for a generated, reopened, verified Office pack."""

    run_id: str
    excel_path: Path
    powerpoint_path: Path
    word_path: Path
    manifest_path: Path
    manifest: Mapping[str, Any]


def load_office_data(
    validation_path: Path,
    claim_manifest_path: Path,
    *,
    artifact_root: Path | None = REPOSITORY_ROOT,
) -> OfficeData:
    """Load, challenge, and normalize the two approved reporting inputs."""

    validation = load_json_object(validation_path)
    claim_manifest = load_json_object(claim_manifest_path)
    build_report_bundle(
        validation,
        validation_path=validation_path,
        artifact_root=artifact_root,
    )

    validation_sha256 = _sha256(validation_path)
    claim_manifest_sha256 = _sha256(claim_manifest_path)
    run_id = _text(validation, "run_id", "$")
    generated_at = _text(validation, "generated_at", "$")
    generated_datetime = _parse_datetime(generated_at)
    packets = _gate_packets(validation)
    gate_1 = _checks(packets[1])
    gate_2 = _checks(packets[2])
    gate_3 = _checks(packets[3])
    findings = _mapping(validation, "verified_findings", "$")

    source_id = _text(_evidence(gate_1, "source_registered"), "source_id", "G1")
    source_sha256 = _text(_evidence(gate_1, "sha256_recorded"), "sha256", "G1")
    if re.fullmatch(r"[0-9a-f]{64}", source_sha256) is None:
        raise OfficePackError("source SHA-256 must contain 64 lowercase hex digits")

    row_check = _evidence(gate_2, "wide_to_long_rows")
    bill_check = _evidence(gate_2, "bill_amounts")
    payment_check = _evidence(gate_2, "payment_amounts")
    test_check = _evidence(gate_2, "sql_test_suite")
    outcome_check = _evidence(gate_3, "outcome_summary")
    definition_version = _text(
        _evidence(gate_3, "definition_version"), "version", "G3"
    )
    signal_count = _positive_int(
        _evidence(gate_3, "signal_count"), "count", "G3"
    )

    source_rows = _positive_int(findings, "source_rows", "findings")
    raw_rows = _positive_int(findings, "raw_rows", "findings")
    borrower_rows = _positive_int(findings, "borrower_rows", "findings")
    account_month_rows = _positive_int(
        findings, "account_month_rows", "findings"
    )
    months = _positive_int(row_check, "months_per_borrower", "G2")
    sql_test_count = _positive_int(findings, "sql_test_count", "findings")
    sql_failed_count = _nonnegative_int(
        findings, "sql_failed_count", "findings"
    )
    _equal(source_rows, raw_rows, "source/raw rows")
    _equal(raw_rows, borrower_rows, "raw/borrower rows")
    _equal(
        account_month_rows,
        borrower_rows * months,
        "borrower-month population",
    )
    _equal(
        sql_test_count,
        _positive_int(test_check, "executed_count", "G2"),
        "SQL test count",
    )
    _equal(
        sql_failed_count,
        _nonnegative_int(test_check, "failed_count", "G2"),
        "SQL failed count",
    )
    if sql_failed_count != 0:
        raise OfficePackError("Office pack requires zero failed SQL tests")

    signal_rows = _outcome_rows(
        findings.get("signal_outcomes"),
        group_field="bucket",
        require_signal=True,
    )
    band_rows = _outcome_rows(
        findings.get("risk_band_outcomes"),
        group_field="risk_band",
        require_signal=False,
    )
    _equal(
        len(signal_rows),
        _positive_int(outcome_check, "summary_row_count", "G3"),
        "signal bucket count",
    )
    _equal(
        len(band_rows),
        _positive_int(outcome_check, "risk_band_row_count", "G3"),
        "risk band count",
    )
    _equal(
        len({row["signal_id"] for row in signal_rows}),
        signal_count,
        "signal definition coverage",
    )
    for signal_id in {str(row["signal_id"]) for row in signal_rows}:
        total = sum(
            int(row["sample_count"])
            for row in signal_rows
            if row["signal_id"] == signal_id
        )
        _equal(total, borrower_rows, f"{signal_id} population")
    _equal(
        sum(int(row["sample_count"]) for row in band_rows),
        borrower_rows,
        "risk-band population",
    )

    claims = _validate_claim_manifest(
        claim_manifest,
        run_id=run_id,
        definition_version=definition_version,
        validation_sha256=validation_sha256,
        generated_at=generated_at,
    )
    gate_rows = _gate_rows(packets)
    dq_rows = _dq_rows(gate_2, gate_3, sql_test_count)
    limitations = (
        IFRS9_LIMITATION,
        "Observed default rates are descriptive associations, not causal effects.",
        (
            "Risk bands are explainable prototype rules, not approved bank "
            "ratings or IFRS 9 Stage assignments."
        ),
        (
            "The independent validation artifact contains the aggregate SQL "
            "test count, not all individual test-result records."
        ),
        (
            "Displayed percentages are rounded to two decimals; exact counts "
            "and exact decimal-rate strings remain the audit authority."
        ),
        (
            "Gate 1–3 machine evidence does not constitute Gate 5 review or "
            "human approval for external publication."
        ),
    )
    return OfficeData(
        run_id=run_id,
        generated_at=generated_at,
        generated_datetime=generated_datetime,
        definition_version=definition_version,
        source_id=source_id,
        source_sha256=source_sha256,
        validation_path=_portable_path(validation_path),
        validation_sha256=validation_sha256,
        claim_manifest_path=_portable_path(claim_manifest_path),
        claim_manifest_sha256=claim_manifest_sha256,
        source_rows=source_rows,
        raw_rows=raw_rows,
        borrower_rows=borrower_rows,
        account_month_rows=account_month_rows,
        months_per_borrower=months,
        bill_source_total=_decimal_text(bill_check.get("source_total"), "bill source"),
        bill_derived_total=_decimal_text(
            bill_check.get("derived_total"), "bill derived"
        ),
        bill_difference=_decimal_text(bill_check.get("difference"), "bill difference"),
        payment_source_total=_decimal_text(
            payment_check.get("source_total"), "payment source"
        ),
        payment_derived_total=_decimal_text(
            payment_check.get("derived_total"), "payment derived"
        ),
        payment_difference=_decimal_text(
            payment_check.get("difference"), "payment difference"
        ),
        sql_test_count=sql_test_count,
        sql_failed_count=sql_failed_count,
        signal_count=signal_count,
        signal_rows=signal_rows,
        band_rows=band_rows,
        gate_rows=gate_rows,
        dq_rows=dq_rows,
        claims=claims,
        limitations=limitations,
    )


def generate_office_pack(
    validation_path: Path,
    claim_manifest_path: Path,
    output_dir: Path,
    *,
    artifact_root: Path | None = REPOSITORY_ROOT,
) -> OfficePackResult:
    """Generate, reopen, reconcile, and publish four new run-bound files."""

    data = load_office_data(
        validation_path,
        claim_manifest_path,
        artifact_root=artifact_root,
    )
    token = _safe_token(data.run_id)
    output_dir.mkdir(parents=True, exist_ok=True)
    final_paths = {
        "excel": output_dir / f"wave1_office_pack__{token}.xlsx",
        "powerpoint": output_dir / f"wave1_office_pack__{token}.pptx",
        "word": output_dir / f"wave1_office_pack__{token}.docx",
        "manifest": output_dir / f"wave1_office_manifest__{token}.json",
    }
    existing = [path for path in final_paths.values() if path.exists()]
    if existing:
        raise OfficePackError(
            "refusing to overwrite existing Office output(s): "
            + ", ".join(str(path) for path in existing)
        )

    with tempfile.TemporaryDirectory(prefix="wave1-office-pack-", dir=output_dir) as tmp:
        staging = Path(tmp)
        staged_paths = {
            key: staging / path.name
            for key, path in final_paths.items()
        }
        _build_excel(data, staged_paths["excel"])
        _build_powerpoint(data, staged_paths["powerpoint"])
        _build_word(data, staged_paths["word"])
        for key in ("excel", "powerpoint", "word"):
            _normalize_ooxml(
                staged_paths[key],
                data.generated_datetime.strftime("%Y-%m-%dT%H:%M:%SZ"),
            )

        excel_receipt = _verify_excel(data, staged_paths["excel"])
        powerpoint_receipt = _verify_powerpoint(
            data, staged_paths["powerpoint"]
        )
        word_receipt = _verify_word(data, staged_paths["word"])
        _verify_cross_artifact(
            data,
            excel_receipt,
            powerpoint_receipt,
            word_receipt,
        )

        artifact_records = []
        for key in ("excel", "powerpoint", "word"):
            path = staged_paths[key]
            artifact_records.append(
                {
                    "artifact_type": key,
                    "path": _portable_path(final_paths[key]),
                    "sha256": _sha256(path),
                    "bytes": path.stat().st_size,
                    "reopened": True,
                    "cross_artifact_status": "pass",
                }
            )
        manifest = _office_manifest(data, artifact_records)
        staged_paths["manifest"].write_text(
            json.dumps(manifest, ensure_ascii=False, indent=2, sort_keys=True)
            + "\n",
            encoding="utf-8",
        )
        reloaded_manifest = load_json_object(staged_paths["manifest"])
        if reloaded_manifest != manifest:
            raise OfficePackError("Office manifest did not round-trip exactly")

        for key in ("excel", "powerpoint", "word", "manifest"):
            _copy_new(staged_paths[key], final_paths[key])

    return OfficePackResult(
        run_id=data.run_id,
        excel_path=final_paths["excel"],
        powerpoint_path=final_paths["powerpoint"],
        word_path=final_paths["word"],
        manifest_path=final_paths["manifest"],
        manifest=manifest,
    )


def _build_excel(data: OfficeData, path: Path) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.creator = "Credit Risk Signal Lab"
    workbook.properties.title = "Wave 1 Internal Audit Workbook"
    workbook.properties.subject = INTERNAL_STATUS
    workbook.properties.created = data.generated_datetime
    workbook.properties.modified = data.generated_datetime

    control = workbook.create_sheet("Control")
    control.merge_cells("A1:D1")
    control["A1"] = "Wave 1 Credit Risk Signal Lab — Internal Audit Workbook"
    control["A1"].font = Font(size=18, bold=True, color=WHITE)
    control["A1"].fill = PatternFill("solid", fgColor=NAVY)
    control["A1"].alignment = Alignment(horizontal="center")
    control.merge_cells("A2:D2")
    control["A2"] = INTERNAL_STATUS
    control["A2"].font = Font(bold=True, color=RED)
    control["A2"].fill = PatternFill("solid", fgColor="FFF2CC")
    control["A2"].alignment = Alignment(horizontal="center")
    control_rows = [
        ("Run ID", data.run_id),
        ("Signal definition version", data.definition_version),
        ("Validation generated at", data.generated_at),
        ("Source ID", data.source_id),
        ("Source SHA-256", data.source_sha256),
        ("Validation artifact", data.validation_path),
        ("Validation SHA-256", data.validation_sha256),
        ("Claim manifest", data.claim_manifest_path),
        ("Claim manifest SHA-256", data.claim_manifest_sha256),
        ("Publication", "pending Gate 5 and human approval"),
        ("Limitation", IFRS9_LIMITATION),
    ]
    for row_index, (label, value) in enumerate(control_rows, start=4):
        control.cell(row_index, 1, label).font = Font(bold=True, color=NAVY)
        control.cell(row_index, 2, value)
        control.merge_cells(
            start_row=row_index,
            start_column=2,
            end_row=row_index,
            end_column=4,
        )
        control.cell(row_index, 2).alignment = Alignment(wrap_text=True)
    control.column_dimensions["A"].width = 28
    control.column_dimensions["B"].width = 34
    control.column_dimensions["C"].width = 20
    control.column_dimensions["D"].width = 20

    reconciliation = workbook.create_sheet("Reconciliation")
    reconciliation.append(
        [
            "Metric",
            "Source / expected",
            "Derived / actual",
            "Difference",
            "Status",
            "Claim ID",
            "Evidence",
        ]
    )
    reconciliation_rows = _reconciliation_rows(data)
    for row in reconciliation_rows:
        reconciliation.append(
            [
                row["metric"],
                row["source"],
                row["derived"],
                row["difference"],
                row["status"],
                row["claim_id"],
                row["evidence"],
            ]
        )
    _format_sheet(reconciliation, freeze="A2", widths=[38, 20, 20, 16, 12, 16, 55])
    _add_table(reconciliation, "ReconciliationTable")

    dq = workbook.create_sheet("DQ_Warnings")
    dq.append(
        [
            "Control ID",
            "Severity",
            "Status",
            "Observed",
            "Threshold / expectation",
            "Affected output",
            "Evidence / warning",
        ]
    )
    for row in data.dq_rows:
        dq.append(
            [
                row["control_id"],
                row["severity"],
                row["status"],
                row["observed"],
                row["expectation"],
                row["affected_output"],
                row["evidence"],
            ]
        )
    _format_sheet(dq, freeze="A2", widths=[28, 13, 15, 18, 26, 30, 62])
    _add_table(dq, "DQWarningsTable")
    for row_index in range(2, dq.max_row + 1):
        status = dq.cell(row_index, 3).value
        if status == "LIMITATION":
            dq.cell(row_index, 3).fill = PatternFill("solid", fgColor="FFF2CC")
        else:
            dq.cell(row_index, 3).fill = PatternFill("solid", fgColor="E2F0D9")

    gates = workbook.create_sheet("Gate_Checks")
    gates.append(["Gate", "Check ID", "Status", "Evidence JSON"])
    for row in data.gate_rows:
        gates.append(
            [
                row["gate"],
                row["check_id"],
                row["status"],
                json.dumps(row["evidence"], ensure_ascii=False, sort_keys=True),
            ]
        )
    _format_sheet(gates, freeze="A2", widths=[9, 32, 13, 110])
    _add_table(gates, "GateChecksTable")

    signals = workbook.create_sheet("Signal_Buckets")
    signals.append(
        [
            "Signal ID",
            "Bucket",
            "Sample count",
            "Default count",
            "Observed default rate exact",
            "Observed default rate",
            "Claim ID",
        ]
    )
    for row in data.signal_rows:
        signals.append(
            [
                row["signal_id"],
                row["bucket"],
                row["sample_count"],
                row["default_count"],
                row["observed_default_rate"],
                float(Decimal(str(row["observed_default_rate"]))),
                "W1-CLM-005",
            ]
        )
    _format_sheet(signals, freeze="A2", widths=[34, 20, 16, 16, 34, 22, 16])
    for cell in signals["F"][1:]:
        cell.number_format = "0.00%"
    _add_table(signals, "SignalBucketsTable")

    bands = workbook.create_sheet("Risk_Bands")
    bands.append(
        [
            "Risk band",
            "Sample count",
            "Default count",
            "Observed default rate exact",
            "Observed default rate",
            "Claim ID",
        ]
    )
    for row in data.band_rows:
        bands.append(
            [
                row["risk_band"],
                row["sample_count"],
                row["default_count"],
                row["observed_default_rate"],
                float(Decimal(str(row["observed_default_rate"]))),
                "W1-CLM-006",
            ]
        )
    _format_sheet(bands, freeze="A2", widths=[18, 16, 16, 34, 22, 16])
    for cell in bands["E"][1:]:
        cell.number_format = "0.00%"
    _add_table(bands, "RiskBandsTable")
    chart = BarChart()
    chart.type = "col"
    chart.style = 10
    chart.title = "Observed default rate by initial risk band"
    chart.y_axis.title = "Observed default rate"
    chart.y_axis.numFmt = "0%"
    chart.x_axis.title = "Risk band"
    chart.height = 8
    chart.width = 14
    chart.add_data(
        Reference(bands, min_col=5, min_row=1, max_row=1 + len(data.band_rows)),
        titles_from_data=True,
    )
    chart.set_categories(
        Reference(bands, min_col=1, min_row=2, max_row=1 + len(data.band_rows))
    )
    bands.add_chart(chart, "H2")

    limitations = workbook.create_sheet("Limitations")
    limitations.append(["Type", "Detail"])
    for item in data.limitations:
        limitations.append(["Limitation", item])
    limitations.append(
        [
            "Publication control",
            INTERNAL_STATUS,
        ]
    )
    _format_sheet(limitations, freeze="A2", widths=[24, 110])
    _add_table(limitations, "LimitationsTable")

    evidence = workbook.create_sheet("Evidence")
    evidence.append(
        [
            "Claim ID",
            "Claim",
            "Run ID",
            "Validation SHA-256",
            "JSON pointers",
        ]
    )
    for claim in data.claims:
        evidence.append(
            [
                claim["claim_id"],
                claim["claim"],
                data.run_id,
                data.validation_sha256,
                " | ".join(str(item) for item in claim["json_pointers"]),
            ]
        )
    _format_sheet(evidence, freeze="A2", widths=[16, 68, 24, 68, 90])
    _add_table(evidence, "EvidenceTable")

    workbook.save(path)


def _build_powerpoint(data: OfficeData, path: Path) -> None:
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    presentation.core_properties.title = "Wave 1 Credit Risk Signal Lab"
    presentation.core_properties.subject = INTERNAL_STATUS
    presentation.core_properties.author = "Credit Risk Signal Lab"
    presentation.core_properties.created = data.generated_datetime
    presentation.core_properties.modified = data.generated_datetime

    slide = _blank_slide(presentation)
    _add_banner(slide)
    _add_text(
        slide,
        0.75,
        1.35,
        11.8,
        1.2,
        "Credit Risk Signal Lab",
        30,
        bold=True,
        color=NAVY,
    )
    _add_text(
        slide,
        0.78,
        2.45,
        11.6,
        0.8,
        "Wave 1 source, SQL QA, and explainable signal evidence",
        20,
        color=DARK,
    )
    _add_text(
        slide,
        0.78,
        3.45,
        11.6,
        1.0,
        (
            f"Run {data.run_id}  |  signal definition {data.definition_version}\n"
            "Prepared for EY FSRM-style internal review"
        ),
        16,
        color=BLUE,
    )
    _add_footer(slide, data, 1)

    slide = _blank_slide(presentation)
    _add_banner(slide)
    _add_title(slide, "Validated Wave 1 evidence is reproducible and reviewable")
    cards = [
        ("Borrowers", _integer(data.borrower_rows)),
        ("Borrower-months", _integer(data.account_month_rows)),
        ("SQL tests", _integer(data.sql_test_count)),
        ("Failed tests", _integer(data.sql_failed_count)),
    ]
    for index, (label, value) in enumerate(cards):
        x = 0.75 + index * 3.1
        _add_card(slide, x, 1.65, 2.75, 1.35, label, value)
    _add_text(
        slide,
        0.8,
        3.45,
        11.7,
        1.6,
        (
            "Gate 1–3 independent checks passed for one run ID. Source/raw, "
            "borrower, six-month borrower-month, bill, and payment populations "
            "reconcile under the approved controls. This supports controlled "
            "analysis; it is not model or publication approval."
        ),
        17,
        color=DARK,
    )
    _add_footer(slide, data, 2)

    slide = _blank_slide(presentation)
    _add_banner(slide)
    _add_title(slide, "The relational transformation preserves rows and amounts")
    rows = [
        ["Metric", "Source / expected", "Derived / actual", "Difference"],
        *[
            [
                str(item["metric"]),
                str(item["source"]),
                str(item["derived"]),
                str(item["difference"]),
            ]
            for item in _reconciliation_rows(data)
        ],
    ]
    _add_ppt_table(slide, rows, 0.75, 1.45, 11.85, 3.55, [4.4, 2.4, 2.4, 1.8])
    _add_text(
        slide,
        0.8,
        5.25,
        11.6,
        0.8,
        (
            "The exact row ties and validated amount differences protect the "
            "downstream borrower-level signal denominator."
        ),
        15,
        color=DARK,
    )
    _add_footer(slide, data, 3)

    slide = _blank_slide(presentation)
    _add_banner(slide)
    _add_title(slide, "Data-quality controls passed; detailed test IDs remain upstream")
    cards = [
        ("Duplicate borrower keys", "0"),
        ("Orphan borrower-months", "0"),
        ("Invalid codes", "0"),
        ("Timing / leakage violations", "0 / 0"),
    ]
    for index, (label, value) in enumerate(cards):
        x = 0.75 + index * 3.1
        _add_card(slide, x, 1.55, 2.75, 1.45, label, value)
    _add_text(
        slide,
        0.8,
        3.45,
        11.7,
        1.7,
        (
            f"{data.sql_test_count} SQL tests are represented in the validated "
            "aggregate, with zero failures. Individual test IDs are not embedded "
            "in this reporting input and therefore are not recreated here; the "
            "SQL test catalog and upstream audit results remain the detail authority."
        ),
        17,
        color=DARK,
    )
    _add_footer(slide, data, 4)

    slide = _blank_slide(presentation)
    _add_banner(slide)
    _add_title(slide, "Observed default rates separate across initial risk bands")
    chart_data = ChartData()
    chart_data.categories = [str(row["risk_band"]) for row in data.band_rows]
    chart_data.add_series(
        "Observed default rate",
        [float(Decimal(str(row["observed_default_rate"]))) for row in data.band_rows],
    )
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(0.7),
        Inches(1.5),
        Inches(7.4),
        Inches(4.7),
        chart_data,
    ).chart
    chart.has_legend = False
    chart.value_axis.tick_labels.number_format = "0%"
    chart.value_axis.maximum_scale = 0.6
    chart.value_axis.minimum_scale = 0
    band_rows = [
        ["Band", "Sample", "Defaults", "Rate"],
        *[
            [
                str(row["risk_band"]),
                _integer(int(row["sample_count"])),
                _integer(int(row["default_count"])),
                _percent(row["observed_default_rate"]),
            ]
            for row in data.band_rows
        ],
    ]
    _add_ppt_table(slide, band_rows, 8.35, 1.7, 4.25, 2.35, [1.05, 1.15, 1.15, 0.9])
    _add_text(
        slide,
        8.4,
        4.35,
        4.05,
        1.25,
        (
            "The comparison is descriptive within this public source. The bands "
            "are prototype review groups—not bank ratings or Stage assignments."
        ),
        13,
        color=DARK,
    )
    _add_footer(slide, data, 5)

    selected_signal = [
        row
        for row in data.signal_rows
        if row["signal_id"] == "delinquent_months_6m"
    ]
    if len(selected_signal) < 2:
        raise OfficePackError(
            "delinquent_months_6m needs at least two buckets for the PPT chart"
        )
    slide = _blank_slide(presentation)
    _add_banner(slide)
    _add_title(slide, "More delinquent months align with higher observed default")
    chart_data = ChartData()
    chart_data.categories = [str(row["bucket"]) for row in selected_signal]
    chart_data.add_series(
        "Observed default rate",
        [
            float(Decimal(str(row["observed_default_rate"])))
            for row in selected_signal
        ],
    )
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(0.7),
        Inches(1.5),
        Inches(7.4),
        Inches(4.7),
        chart_data,
    ).chart
    chart.has_legend = False
    chart.value_axis.tick_labels.number_format = "0%"
    chart.value_axis.maximum_scale = 0.6
    chart.value_axis.minimum_scale = 0
    signal_rows = [
        ["6m delinquent months", "Sample", "Defaults", "Rate"],
        *[
            [
                str(row["bucket"]),
                _integer(int(row["sample_count"])),
                _integer(int(row["default_count"])),
                _percent(row["observed_default_rate"]),
            ]
            for row in selected_signal
        ],
    ]
    _add_ppt_table(
        slide,
        signal_rows,
        8.25,
        1.65,
        4.35,
        2.45,
        [1.45, 1.05, 1.05, 0.8],
    )
    _add_text(
        slide,
        8.35,
        4.35,
        4.0,
        1.25,
        (
            "All 23 validated signal buckets are preserved in the Excel and "
            "Word artifacts. This chart shows one economically interpretable "
            "path and does not establish causality or out-of-sample performance."
        ),
        13,
        color=DARK,
    )
    _add_footer(slide, data, 6)

    slide = _blank_slide(presentation)
    _add_banner(slide)
    _add_title(slide, "Traceability is complete; publication approval is not")
    _add_text(
        slide,
        0.75,
        1.45,
        5.8,
        3.6,
        (
            "Evidence path\n\n"
            "Claim ID → Office manifest → independent validation JSON → "
            "Gate checks → deterministic SQL/code → raw source + SHA-256\n\n"
            f"Source SHA-256\n{data.source_sha256}"
        ),
        15,
        color=DARK,
    )
    _add_text(
        slide,
        6.9,
        1.45,
        5.6,
        3.9,
        (
            "Required next controls\n\n"
            "1. Reconcile Excel, PowerPoint, Word, and report hashes/values.\n"
            "2. Complete independent Gate 5 review.\n"
            "3. Obtain explicit human approval before external use.\n\n"
            f"Limitation: {IFRS9_LIMITATION}"
        ),
        15,
        color=DARK,
    )
    _add_footer(slide, data, 7)
    presentation.save(path)


def _build_word(data: OfficeData, path: Path) -> None:
    document = Document()
    document.core_properties.title = "Wave 1 Technical Methodology and Test Report"
    document.core_properties.subject = INTERNAL_STATUS
    document.core_properties.author = "Credit Risk Signal Lab"
    document.core_properties.created = data.generated_datetime
    document.core_properties.modified = data.generated_datetime
    section = document.sections[0]
    section.top_margin = DocxInches(0.7)
    section.bottom_margin = DocxInches(0.7)
    section.left_margin = DocxInches(0.75)
    section.right_margin = DocxInches(0.75)
    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = DocxPt(9)
    styles["Title"].font.name = "Arial"
    styles["Title"].font.size = DocxPt(24)
    styles["Heading 1"].font.name = "Arial"
    styles["Heading 1"].font.color.rgb = _docx_rgb(NAVY)
    styles["Heading 2"].font.name = "Arial"
    styles["Heading 2"].font.color.rgb = _docx_rgb(BLUE)

    title = document.add_heading(
        "Wave 1 Technical Methodology and Test Report",
        0,
    )
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    status = document.add_paragraph()
    status.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = status.add_run(INTERNAL_STATUS)
    run.bold = True
    run.font.color.rgb = _docx_rgb(RED)

    metadata = [
        ["Run ID", data.run_id],
        ["Signal definition version", data.definition_version],
        ["Validation generated at", data.generated_at],
        ["Source ID", data.source_id],
        ["Source SHA-256", data.source_sha256],
        ["Publication", "Pending Gate 5 and human approval"],
    ]
    _add_docx_table(document, [["Field", "Value"], *metadata], [1.8, 5.8])

    document.add_heading("Technical summary", level=1)
    document.add_paragraph(
        (
            f"Independent Gate 1–3 evidence for run {data.run_id} validates "
            f"{_integer(data.borrower_rows)} borrowers, "
            f"{_integer(data.account_month_rows)} borrower-month rows, and "
            f"{data.sql_test_count} SQL tests with {data.sql_failed_count} "
            "failures. Row and amount reconciliations satisfy the approved "
            "controls. Observed rates are descriptive and internal-only."
        )
    )

    document.add_heading("Reconciled populations protect the signal denominator", level=1)
    reconciliation_rows = [
        ["Metric", "Source / expected", "Derived / actual", "Difference", "Status"],
        *[
            [
                str(row["metric"]),
                str(row["source"]),
                str(row["derived"]),
                str(row["difference"]),
                str(row["status"]),
            ]
            for row in _reconciliation_rows(data)
        ],
    ]
    _add_docx_table(document, reconciliation_rows, [2.6, 1.25, 1.25, 1.0, 0.8])
    document.add_paragraph(
        (
            "These checks establish transformation consistency. They do not "
            "establish causal, predictive, calibration, or production-model validity."
        )
    )

    document.add_heading("Data-quality controls passed with one reporting limitation", level=1)
    dq_rows = [
        ["Control", "Status", "Observed", "Expectation", "Evidence / warning"],
        *[
            [
                str(row["control_id"]),
                str(row["status"]),
                str(row["observed"]),
                str(row["expectation"]),
                str(row["evidence"]),
            ]
            for row in data.dq_rows
        ],
    ]
    _add_docx_table(document, dq_rows, [1.7, 0.8, 1.0, 1.3, 3.0])

    document.add_heading("Observed default rates differ by initial risk band", level=1)
    document.add_paragraph(
        (
            "Counts remain the denominator authority. Displayed percentages "
            "are rounded to two decimal places."
        )
    )
    band_rows = [
        ["Risk band", "Sample", "Defaults", "Exact rate", "Display rate"],
        *[
            [
                str(row["risk_band"]),
                str(row["sample_count"]),
                str(row["default_count"]),
                str(row["observed_default_rate"]),
                _percent(row["observed_default_rate"]),
            ]
            for row in data.band_rows
        ],
    ]
    _add_docx_table(document, band_rows, [1.0, 1.0, 1.0, 2.8, 1.0])

    document.add_heading("All validated signal-bucket outcomes remain auditable", level=1)
    document.add_paragraph(
        (
            f"The {len(data.signal_rows)} rows below cover {data.signal_count} "
            "versioned signals. Each signal's bucket samples reconcile to the "
            "borrower population."
        )
    )
    signal_rows = [
        ["Signal ID", "Bucket", "Sample", "Defaults", "Exact rate", "Display rate"],
        *[
            [
                str(row["signal_id"]),
                str(row["bucket"]),
                str(row["sample_count"]),
                str(row["default_count"]),
                str(row["observed_default_rate"]),
                _percent(row["observed_default_rate"]),
            ]
            for row in data.signal_rows
        ],
    ]
    _add_docx_table(document, signal_rows, [1.8, 1.2, 0.8, 0.8, 2.35, 0.85])

    document.add_heading("Scope, data, and metric definitions", level=1)
    definitions = [
        (
            "Observation grain",
            "Borrower and borrower-month; six source history months per borrower.",
        ),
        (
            "Outcome",
            "Source-provided next-month default label.",
        ),
        (
            "Observed default rate",
            "Default count divided by sample count within each validated bucket.",
        ),
        (
            "Risk band",
            "Explainable prototype grouping under definition version "
            f"{data.definition_version}; not a bank rating or Stage.",
        ),
        (
            "Claim type",
            "Descriptive observed association; neither causal inference nor "
            "out-of-sample model performance.",
        ),
    ]
    _add_docx_table(
        document,
        [["Definition", "Approved report meaning"], *[list(row) for row in definitions]],
        [1.8, 5.8],
    )

    document.add_heading("Methodology and reporting controls", level=1)
    for item in (
        "Validate Gate 1–3 packet structure, referenced artifacts, run identity, and arithmetic.",
        "Match the claim manifest run/version and validation SHA-256.",
        "Read counts and exact decimal rates only from verified_findings and Gate evidence.",
        "Generate all Office outputs from the same in-memory OfficeData contract.",
        "Normalize OOXML package timestamps for deterministic file hashes.",
        "Reopen XLSX/PPTX/DOCX and compare metadata, populations, outcomes, and limitations.",
        "Write a file-hash/value manifest only after cross-artifact checks pass.",
        "Refuse to overwrite any prior Office artifact or manifest.",
    ):
        document.add_paragraph(item, style="List Bullet")

    document.add_heading("Limitations, uncertainty, and robustness checks", level=1)
    for item in data.limitations:
        document.add_paragraph(item, style="List Bullet")

    document.add_heading("Claim-to-evidence mapping", level=1)
    claim_rows = [
        ["Claim ID", "Claim", "JSON evidence pointers"],
        *[
            [
                str(claim["claim_id"]),
                str(claim["claim"]),
                " | ".join(str(item) for item in claim["json_pointers"]),
            ]
            for claim in data.claims
        ],
    ]
    _add_docx_table(document, claim_rows, [1.0, 3.0, 3.7])

    document.add_heading("Recommended next steps", level=1)
    for item in (
        "Review the Office manifest and reconcile its hashes and common values.",
        "Complete independent Gate 5 cross-artifact review.",
        "Obtain explicit human approval before any external portfolio quotation.",
        "In Wave 2, test time-aware stability and calibration before model conclusions.",
    ):
        document.add_paragraph(item, style="List Number")

    document.add_heading("Further questions", level=1)
    document.add_paragraph(
        "Do signal directions persist under time-aware and segment validation?"
    )
    document.add_paragraph(
        "Which public-data limitations prevent each prototype from representing "
        "a bank production process?"
    )

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.text = (
        f"{data.run_id} | definition {data.definition_version} | {INTERNAL_STATUS}"
    )
    document.save(path)


def _verify_excel(data: OfficeData, path: Path) -> Mapping[str, Any]:
    workbook = load_workbook(path, data_only=False, read_only=False)
    required = {
        "Control",
        "Reconciliation",
        "DQ_Warnings",
        "Gate_Checks",
        "Signal_Buckets",
        "Risk_Bands",
        "Limitations",
        "Evidence",
    }
    if set(workbook.sheetnames) != required:
        raise OfficePackError("Excel sheet set does not match the reporting contract")
    control = {
        workbook["Control"].cell(row, 1).value: workbook["Control"].cell(row, 2).value
        for row in range(4, workbook["Control"].max_row + 1)
    }
    _equal(control.get("Run ID"), data.run_id, "Excel run ID")
    _equal(
        control.get("Signal definition version"),
        data.definition_version,
        "Excel definition version",
    )
    _equal(control.get("Source SHA-256"), data.source_sha256, "Excel source hash")
    if INTERNAL_STATUS not in str(workbook["Control"]["A2"].value):
        raise OfficePackError("Excel publication banner is missing")
    if IFRS9_LIMITATION not in str(control.get("Limitation")):
        raise OfficePackError("Excel IFRS 9 limitation is missing")

    reconciliation_rows = _excel_rows(
        workbook["Reconciliation"],
        [
            "Metric",
            "Source / expected",
            "Derived / actual",
            "Difference",
            "Status",
            "Claim ID",
            "Evidence",
        ],
    )
    normalized_reconciliation = [
        {
            "metric": row["Metric"],
            "source": row["Source / expected"],
            "derived": row["Derived / actual"],
            "difference": row["Difference"],
            "status": row["Status"],
            "claim_id": row["Claim ID"],
            "evidence": row["Evidence"],
        }
        for row in reconciliation_rows
    ]
    _equal(
        normalized_reconciliation,
        list(_reconciliation_rows(data)),
        "Excel reconciliation rows",
    )

    dq_rows = _excel_rows(
        workbook["DQ_Warnings"],
        [
            "Control ID",
            "Severity",
            "Status",
            "Observed",
            "Threshold / expectation",
            "Affected output",
            "Evidence / warning",
        ],
    )
    normalized_dq = [
        {
            "control_id": row["Control ID"],
            "severity": row["Severity"],
            "status": row["Status"],
            "observed": row["Observed"],
            "expectation": row["Threshold / expectation"],
            "affected_output": row["Affected output"],
            "evidence": row["Evidence / warning"],
        }
        for row in dq_rows
    ]
    _equal(normalized_dq, list(data.dq_rows), "Excel DQ rows")

    gate_rows = _excel_rows(
        workbook["Gate_Checks"],
        ["Gate", "Check ID", "Status", "Evidence JSON"],
    )
    expected_gate_rows = [
        {
            "Gate": row["gate"],
            "Check ID": row["check_id"],
            "Status": row["status"],
            "Evidence JSON": json.dumps(
                row["evidence"],
                ensure_ascii=False,
                sort_keys=True,
            ),
        }
        for row in data.gate_rows
    ]
    _equal(gate_rows, expected_gate_rows, "Excel Gate rows")

    signal_rows = _excel_rows(
        workbook["Signal_Buckets"],
        ["Signal ID", "Bucket", "Sample count", "Default count", "Observed default rate exact"],
    )
    normalized_signals = [
        {
            "signal_id": row["Signal ID"],
            "bucket": row["Bucket"],
            "sample_count": row["Sample count"],
            "default_count": row["Default count"],
            "observed_default_rate": row["Observed default rate exact"],
        }
        for row in signal_rows
    ]
    _equal(normalized_signals, list(data.signal_rows), "Excel signal rows")

    band_rows = _excel_rows(
        workbook["Risk_Bands"],
        ["Risk band", "Sample count", "Default count", "Observed default rate exact"],
    )
    normalized_bands = [
        {
            "signal_id": None,
            "risk_band": row["Risk band"],
            "sample_count": row["Sample count"],
            "default_count": row["Default count"],
            "observed_default_rate": row["Observed default rate exact"],
        }
        for row in band_rows
    ]
    _equal(normalized_bands, list(data.band_rows), "Excel risk-band rows")
    if len(workbook["Risk_Bands"]._charts) != 1:
        raise OfficePackError("Excel risk-band chart is missing")

    evidence_rows = _excel_rows(
        workbook["Evidence"],
        ["Claim ID", "Claim", "Run ID", "Validation SHA-256", "JSON pointers"],
    )
    expected_evidence = [
        {
            "Claim ID": claim["claim_id"],
            "Claim": claim["claim"],
            "Run ID": data.run_id,
            "Validation SHA-256": data.validation_sha256,
            "JSON pointers": " | ".join(
                str(item) for item in claim["json_pointers"]
            ),
        }
        for claim in data.claims
    ]
    _equal(evidence_rows, expected_evidence, "Excel evidence rows")
    workbook.close()
    return {
        "run_id": data.run_id,
        "definition_version": data.definition_version,
        "source_sha256": data.source_sha256,
        "source_rows": data.source_rows,
        "borrower_rows": data.borrower_rows,
        "account_month_rows": data.account_month_rows,
        "sql_test_count": data.sql_test_count,
        "sql_failed_count": data.sql_failed_count,
        "reconciliation_rows": normalized_reconciliation,
        "signal_rows": normalized_signals,
        "band_rows": normalized_bands,
        "status": INTERNAL_STATUS,
        "limitation": IFRS9_LIMITATION,
    }


def _verify_powerpoint(data: OfficeData, path: Path) -> Mapping[str, Any]:
    presentation = Presentation(path)
    texts: list[str] = []
    chart_count = 0
    tables: list[list[list[str]]] = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                texts.append(shape.text)
            if getattr(shape, "has_table", False):
                table_rows = [
                    [cell.text for cell in row.cells]
                    for row in shape.table.rows
                ]
                tables.append(table_rows)
            if getattr(shape, "has_chart", False):
                chart_count += 1
    joined = "\n".join(texts)
    for expected in (
        data.run_id,
        data.definition_version,
        data.source_sha256,
        INTERNAL_STATUS,
        IFRS9_LIMITATION,
        _integer(data.borrower_rows),
        _integer(data.account_month_rows),
        str(data.sql_test_count),
    ):
        if expected not in joined:
            raise OfficePackError(f"PowerPoint is missing required value: {expected}")
    if chart_count != 2:
        raise OfficePackError("PowerPoint must contain exactly two verified charts")

    reconciliation_table = _ppt_table(tables, "Metric")
    actual_reconciliation = reconciliation_table[1:]
    expected_reconciliation = [
        [
            str(item["metric"]),
            str(item["source"]),
            str(item["derived"]),
            str(item["difference"]),
        ]
        for item in _reconciliation_rows(data)
    ]
    _equal(
        actual_reconciliation,
        expected_reconciliation,
        "PowerPoint reconciliation table",
    )

    band_table = _ppt_table(tables, "Band")
    actual_bands = [
        {
            "risk_band": row[0],
            "sample_count": row[1],
            "default_count": row[2],
            "rate": row[3],
        }
        for row in band_table[1:]
    ]
    expected_bands = [
        {
            "risk_band": str(row["risk_band"]),
            "sample_count": _integer(int(row["sample_count"])),
            "default_count": _integer(int(row["default_count"])),
            "rate": _percent(row["observed_default_rate"]),
        }
        for row in data.band_rows
    ]
    _equal(actual_bands, expected_bands, "PowerPoint risk-band table")

    signal_table = _ppt_table(tables, "6m delinquent months")
    selected = [
        row
        for row in data.signal_rows
        if row["signal_id"] == "delinquent_months_6m"
    ]
    actual_signal = [
        {
            "bucket": row[0],
            "sample_count": row[1],
            "default_count": row[2],
            "rate": row[3],
        }
        for row in signal_table[1:]
    ]
    expected_signal = [
        {
            "bucket": str(row["bucket"]),
            "sample_count": _integer(int(row["sample_count"])),
            "default_count": _integer(int(row["default_count"])),
            "rate": _percent(row["observed_default_rate"]),
        }
        for row in selected
    ]
    _equal(actual_signal, expected_signal, "PowerPoint selected signal table")
    return {
        "run_id": data.run_id,
        "definition_version": data.definition_version,
        "source_sha256": data.source_sha256,
        "source_rows": data.source_rows,
        "borrower_rows": data.borrower_rows,
        "account_month_rows": data.account_month_rows,
        "sql_test_count": data.sql_test_count,
        "sql_failed_count": data.sql_failed_count,
        "reconciliation_rows": list(_reconciliation_rows(data)),
        "band_rows": list(data.band_rows),
        "selected_signal_rows": list(selected),
        "status": INTERNAL_STATUS,
        "limitation": IFRS9_LIMITATION,
        "chart_count": chart_count,
    }


def _verify_word(data: OfficeData, path: Path) -> Mapping[str, Any]:
    document = Document(path)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    table_text = "\n".join(
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    )
    joined = text + "\n" + table_text
    for expected in (
        data.run_id,
        data.definition_version,
        data.source_sha256,
        INTERNAL_STATUS,
        IFRS9_LIMITATION,
        _integer(data.borrower_rows),
        _integer(data.account_month_rows),
        str(data.sql_test_count),
    ):
        if expected not in joined:
            raise OfficePackError(f"Word report is missing required value: {expected}")

    reconciliation_table = _docx_table(document, "Metric")
    reconciliation_headers = [
        cell.text for cell in reconciliation_table.rows[0].cells
    ]
    reconciliation_rows = [
        dict(
            zip(
                reconciliation_headers,
                [cell.text for cell in row.cells],
                strict=True,
            )
        )
        for row in reconciliation_table.rows[1:]
    ]
    normalized_reconciliation = [
        {
            "metric": row["Metric"],
            "source": row["Source / expected"],
            "derived": row["Derived / actual"],
            "difference": row["Difference"],
            "status": row["Status"],
        }
        for row in reconciliation_rows
    ]
    expected_reconciliation = [
        {
            "metric": str(row["metric"]),
            "source": str(row["source"]),
            "derived": str(row["derived"]),
            "difference": str(row["difference"]),
            "status": str(row["status"]),
        }
        for row in _reconciliation_rows(data)
    ]
    _equal(
        normalized_reconciliation,
        expected_reconciliation,
        "Word reconciliation rows",
    )

    dq_table = _docx_table(document, "Control")
    dq_headers = [cell.text for cell in dq_table.rows[0].cells]
    dq_rows = [
        dict(zip(dq_headers, [cell.text for cell in row.cells], strict=True))
        for row in dq_table.rows[1:]
    ]
    normalized_dq = [
        {
            "control_id": row["Control"],
            "status": row["Status"],
            "observed": row["Observed"],
            "expectation": row["Expectation"],
            "evidence": row["Evidence / warning"],
        }
        for row in dq_rows
    ]
    expected_dq = [
        {
            "control_id": str(row["control_id"]),
            "status": str(row["status"]),
            "observed": str(row["observed"]),
            "expectation": str(row["expectation"]),
            "evidence": str(row["evidence"]),
        }
        for row in data.dq_rows
    ]
    _equal(normalized_dq, expected_dq, "Word DQ rows")

    signal_table = _docx_table(document, "Signal ID")
    signal_headers = [cell.text for cell in signal_table.rows[0].cells]
    signal_rows = [
        dict(zip(signal_headers, [cell.text for cell in row.cells], strict=True))
        for row in signal_table.rows[1:]
    ]
    normalized_signals = [
        {
            "signal_id": row["Signal ID"],
            "bucket": row["Bucket"],
            "sample_count": int(row["Sample"]),
            "default_count": int(row["Defaults"]),
            "observed_default_rate": row["Exact rate"],
        }
        for row in signal_rows
    ]
    _equal(normalized_signals, list(data.signal_rows), "Word signal rows")

    band_table = _docx_table(document, "Risk band")
    band_headers = [cell.text for cell in band_table.rows[0].cells]
    band_rows = [
        dict(zip(band_headers, [cell.text for cell in row.cells], strict=True))
        for row in band_table.rows[1:]
    ]
    normalized_bands = [
        {
            "signal_id": None,
            "risk_band": row["Risk band"],
            "sample_count": int(row["Sample"]),
            "default_count": int(row["Defaults"]),
            "observed_default_rate": row["Exact rate"],
        }
        for row in band_rows
    ]
    _equal(normalized_bands, list(data.band_rows), "Word risk-band rows")

    claim_table = _docx_table(document, "Claim ID")
    claim_headers = [cell.text for cell in claim_table.rows[0].cells]
    claim_rows = [
        dict(zip(claim_headers, [cell.text for cell in row.cells], strict=True))
        for row in claim_table.rows[1:]
    ]
    expected_claims = [
        {
            "Claim ID": str(claim["claim_id"]),
            "Claim": str(claim["claim"]),
            "JSON evidence pointers": " | ".join(
                str(item) for item in claim["json_pointers"]
            ),
        }
        for claim in data.claims
    ]
    _equal(claim_rows, expected_claims, "Word claim rows")
    return {
        "run_id": data.run_id,
        "definition_version": data.definition_version,
        "source_sha256": data.source_sha256,
        "source_rows": data.source_rows,
        "borrower_rows": data.borrower_rows,
        "account_month_rows": data.account_month_rows,
        "sql_test_count": data.sql_test_count,
        "sql_failed_count": data.sql_failed_count,
        "reconciliation_rows": normalized_reconciliation,
        "signal_rows": normalized_signals,
        "band_rows": normalized_bands,
        "status": INTERNAL_STATUS,
        "limitation": IFRS9_LIMITATION,
    }


def _verify_cross_artifact(
    data: OfficeData,
    excel: Mapping[str, Any],
    powerpoint: Mapping[str, Any],
    word: Mapping[str, Any],
) -> None:
    common_fields = (
        "run_id",
        "definition_version",
        "source_sha256",
        "source_rows",
        "borrower_rows",
        "account_month_rows",
        "sql_test_count",
        "sql_failed_count",
        "status",
        "limitation",
    )
    for field in common_fields:
        _equal(excel[field], powerpoint[field], f"Excel/PPT {field}")
        _equal(excel[field], word[field], f"Excel/Word {field}")
    _equal(excel["signal_rows"], word["signal_rows"], "Excel/Word signal rows")
    _equal(excel["band_rows"], word["band_rows"], "Excel/Word risk-band rows")
    _equal(excel["band_rows"], powerpoint["band_rows"], "Excel/PPT risk-band rows")
    _equal(
        excel["reconciliation_rows"],
        powerpoint["reconciliation_rows"],
        "Excel/PPT reconciliation rows",
    )
    excel_reconciliation = [
        {
            key: str(value)
            for key, value in row.items()
            if key in {"metric", "source", "derived", "difference", "status"}
        }
        for row in excel["reconciliation_rows"]
    ]
    _equal(
        excel_reconciliation,
        word["reconciliation_rows"],
        "Excel/Word reconciliation rows",
    )
    _equal(len(excel["signal_rows"]), len(data.signal_rows), "Office signal coverage")


def _office_manifest(
    data: OfficeData,
    artifacts: Sequence[Mapping[str, Any]],
) -> dict[str, Any]:
    return {
        "schema_version": "1.0",
        "pack_type": "wave1_internal_office_pack",
        "run_id": data.run_id,
        "signal_definition_version": data.definition_version,
        "validation_generated_at": data.generated_at,
        "publication": {
            "status": "pending",
            "required_gate": 5,
            "human_approval_required": True,
            "label": INTERNAL_STATUS,
        },
        "inputs": [
            {
                "input_type": "independent_validation",
                "path": data.validation_path,
                "sha256": data.validation_sha256,
            },
            {
                "input_type": "claim_manifest",
                "path": data.claim_manifest_path,
                "sha256": data.claim_manifest_sha256,
            },
        ],
        "source": {
            "source_id": data.source_id,
            "sha256": data.source_sha256,
        },
        "artifacts": list(artifacts),
        "cross_artifact_values": {
            "source_rows": data.source_rows,
            "raw_rows": data.raw_rows,
            "borrower_rows": data.borrower_rows,
            "account_month_rows": data.account_month_rows,
            "months_per_borrower": data.months_per_borrower,
            "bill_difference": data.bill_difference,
            "payment_difference": data.payment_difference,
            "sql_test_count": data.sql_test_count,
            "sql_failed_count": data.sql_failed_count,
            "signal_count": data.signal_count,
            "signal_bucket_count": len(data.signal_rows),
            "risk_band_count": len(data.band_rows),
            "signal_rows": list(data.signal_rows),
            "risk_band_rows": list(data.band_rows),
        },
        "claims": [str(claim["claim_id"]) for claim in data.claims],
        "limitations": list(data.limitations),
        "verification": {
            "status": "pass",
            "xlsx_reopened_with": "openpyxl",
            "pptx_reopened_with": "python-pptx",
            "docx_reopened_with": "python-docx",
            "cross_artifact_values_reconciled": True,
            "binary_packages_timestamp_normalized": True,
        },
        "generator": "src/reporting/build_wave1_office_pack.py",
    }


def _validate_claim_manifest(
    manifest: Mapping[str, Any],
    *,
    run_id: str,
    definition_version: str,
    validation_sha256: str,
    generated_at: str,
) -> tuple[Mapping[str, Any], ...]:
    if manifest.get("schema_version") != "1.0":
        raise OfficePackError("claim manifest schema_version must equal '1.0'")
    _equal(manifest.get("run_id"), run_id, "claim manifest run ID")
    _equal(
        manifest.get("signal_definition_version"),
        definition_version,
        "claim manifest definition version",
    )
    _equal(
        manifest.get("validation_generated_at"),
        generated_at,
        "claim manifest validation timestamp",
    )
    publication = _mapping(manifest, "publication", "claim manifest")
    if (
        publication.get("status") != "pending"
        or publication.get("required_gate") != 5
        or publication.get("human_approval_required") is not True
    ):
        raise OfficePackError("claim manifest publication controls are not pending")
    claims_value = manifest.get("claims")
    if not isinstance(claims_value, list):
        raise OfficePackError("claim manifest claims must be a list")
    claims: list[Mapping[str, Any]] = []
    seen: set[str] = set()
    for index, claim in enumerate(claims_value):
        if not isinstance(claim, Mapping):
            raise OfficePackError(f"claim manifest claim {index} must be an object")
        claim_id = _text(claim, "claim_id", f"claim[{index}]")
        if claim_id in seen:
            raise OfficePackError(f"duplicate claim ID {claim_id}")
        seen.add(claim_id)
        _equal(claim.get("run_id"), run_id, f"{claim_id} run ID")
        _equal(
            claim.get("validation_artifact_sha256"),
            validation_sha256,
            f"{claim_id} validation hash",
        )
        if claim.get("required_gate_status") != "pass":
            raise OfficePackError(f"{claim_id} required gate status is not pass")
        _text(claim, "claim", claim_id)
        pointers = claim.get("json_pointers")
        if not isinstance(pointers, list) or not pointers or not all(
            isinstance(item, str) and item.strip() for item in pointers
        ):
            raise OfficePackError(f"{claim_id} JSON pointers are incomplete")
        claims.append(claim)
    _equal(seen, set(REQUIRED_CLAIM_IDS), "claim ID coverage")
    return tuple(sorted(claims, key=lambda item: str(item["claim_id"])))


def _gate_packets(validation: Mapping[str, Any]) -> dict[int, Mapping[str, Any]]:
    value = validation.get("gate_packets")
    if not isinstance(value, list):
        raise OfficePackError("gate_packets must be a list")
    result: dict[int, Mapping[str, Any]] = {}
    for packet in value:
        if not isinstance(packet, Mapping):
            raise OfficePackError("each Gate packet must be an object")
        gate = packet.get("gate")
        if not isinstance(gate, int) or isinstance(gate, bool):
            raise OfficePackError("Gate packet number must be an integer")
        result[gate] = packet
    if set(result) != {1, 2, 3}:
        raise OfficePackError("Office pack requires Gate 1–3 packets")
    return result


def _checks(packet: Mapping[str, Any]) -> dict[str, Mapping[str, Any]]:
    value = packet.get("checks")
    if not isinstance(value, list):
        raise OfficePackError("Gate checks must be a list")
    result: dict[str, Mapping[str, Any]] = {}
    for check in value:
        if not isinstance(check, Mapping):
            raise OfficePackError("each Gate check must be an object")
        check_id = _text(check, "check_id", "Gate check")
        result[check_id] = check
    return result


def _evidence(
    checks: Mapping[str, Mapping[str, Any]],
    check_id: str,
) -> Mapping[str, Any]:
    check = checks.get(check_id)
    if check is None:
        raise OfficePackError(f"missing Gate check {check_id}")
    if check.get("status") != "pass":
        raise OfficePackError(f"Gate check {check_id} is not pass")
    return _mapping(check, "evidence", check_id)


def _gate_rows(
    packets: Mapping[int, Mapping[str, Any]],
) -> tuple[Mapping[str, Any], ...]:
    rows: list[Mapping[str, Any]] = []
    for gate in sorted(packets):
        checks = packets[gate].get("checks")
        if not isinstance(checks, list):
            raise OfficePackError("Gate checks must be a list")
        for check in checks:
            if not isinstance(check, Mapping):
                raise OfficePackError("Gate check must be an object")
            rows.append(
                {
                    "gate": gate,
                    "check_id": _text(check, "check_id", "Gate check"),
                    "status": _text(check, "status", "Gate check"),
                    "evidence": _mapping(check, "evidence", "Gate check"),
                }
            )
    return tuple(rows)


def _dq_rows(
    gate_2: Mapping[str, Mapping[str, Any]],
    gate_3: Mapping[str, Mapping[str, Any]],
    sql_test_count: int,
) -> tuple[Mapping[str, Any], ...]:
    primary = _evidence(gate_2, "primary_key")
    foreign = _evidence(gate_2, "foreign_key")
    domains = _evidence(gate_2, "code_domains")
    suite = _evidence(gate_2, "sql_test_suite")
    timing = _evidence(gate_3, "observation_timing")
    leakage = _evidence(gate_3, "leakage")
    rows = [
        ("DQ-PK-BORROWER", "Critical", "PASS", primary["duplicate_count"], "0", "core.borrower", "Gate 2 primary_key"),
        ("DQ-PK-ACCOUNT-MONTH", "Critical", "PASS", primary["account_month_duplicate_count"], "0", "core.account_month", "Gate 2 primary_key"),
        ("DQ-NULL-BORROWER-KEY", "Critical", "PASS", primary["null_key_count"], "0", "core.borrower", "Gate 2 primary_key"),
        ("DQ-FK-ACCOUNT-MONTH", "Critical", "PASS", foreign["orphan_count"], "0", "core.account_month", "Gate 2 foreign_key"),
        ("DQ-CODE-DOMAIN", "High", "PASS", domains["invalid_count"], "0", "core tables", "Gate 2 code_domains"),
        ("DQ-MISSING-HISTORY", "High", "PASS", domains["missing_required_count"], "0", "core.account_month", "Gate 2 code_domains"),
        ("DQ-SQL-FAILURES", "Critical", "PASS", suite["failed_count"], "0", "all Wave 1 marts", "Gate 2 sql_test_suite"),
        ("DQ-TIMING", "Critical", "PASS", timing["violation_count"], "0", "risk signals", "Gate 3 observation_timing"),
        ("DQ-LEAKAGE", "Critical", "PASS", leakage["violation_count"], "0", "risk signals", "Gate 3 leakage"),
        (
            "DQ-SQL-DETAIL-COVERAGE",
            "Info",
            "LIMITATION",
            f"{sql_test_count} aggregate tests",
            "individual test IDs remain upstream",
            "test-level audit drilldown",
            "validation input exposes counts but not individual test records",
        ),
    ]
    return tuple(
        {
            "control_id": row[0],
            "severity": row[1],
            "status": row[2],
            "observed": row[3],
            "expectation": row[4],
            "affected_output": row[5],
            "evidence": row[6],
        }
        for row in rows
    )


def _outcome_rows(
    value: object,
    *,
    group_field: str,
    require_signal: bool,
) -> tuple[Mapping[str, Any], ...]:
    if not isinstance(value, list) or not value:
        raise OfficePackError(f"{group_field} rows must be a non-empty list")
    rows: list[Mapping[str, Any]] = []
    seen: set[tuple[str, str]] = set()
    for index, row in enumerate(value):
        if not isinstance(row, Mapping):
            raise OfficePackError(f"{group_field} row {index} must be an object")
        signal_id = row.get("signal_id")
        if require_signal and (not isinstance(signal_id, str) or not signal_id):
            raise OfficePackError(f"{group_field} row {index} requires signal_id")
        if not require_signal and signal_id is not None:
            raise OfficePackError("risk-band signal_id must be null")
        group = _text(row, group_field, f"{group_field}[{index}]")
        key = (str(signal_id), group)
        if key in seen:
            raise OfficePackError(f"duplicate outcome row {key}")
        seen.add(key)
        sample = _positive_int(row, "sample_count", f"{group_field}[{index}]")
        defaults = _nonnegative_int(
            row, "default_count", f"{group_field}[{index}]"
        )
        if defaults > sample:
            raise OfficePackError("default count cannot exceed sample count")
        rate = _decimal_text(
            row.get("observed_default_rate"),
            f"{group_field}[{index}].observed_default_rate",
        )
        if Decimal(defaults) / Decimal(sample) != Decimal(rate):
            raise OfficePackError(f"{group_field} row {index} rate is inconsistent")
        normalized: dict[str, Any] = {
            "signal_id": signal_id,
            group_field: group,
            "sample_count": sample,
            "default_count": defaults,
            "observed_default_rate": rate,
        }
        rows.append(normalized)
    return tuple(rows)


def _reconciliation_rows(data: OfficeData) -> tuple[Mapping[str, Any], ...]:
    return (
        {
            "metric": "Source rows → raw rows",
            "source": data.source_rows,
            "derived": data.raw_rows,
            "difference": data.raw_rows - data.source_rows,
            "status": "PASS",
            "claim_id": "W1-CLM-001",
            "evidence": "Gate 1 raw_rows_reconciled",
        },
        {
            "metric": "Raw rows → borrowers",
            "source": data.raw_rows,
            "derived": data.borrower_rows,
            "difference": data.borrower_rows - data.raw_rows,
            "status": "PASS",
            "claim_id": "W1-CLM-002",
            "evidence": "Gate 2 wide_to_long_rows",
        },
        {
            "metric": f"Borrowers × {data.months_per_borrower} → borrower-months",
            "source": data.borrower_rows * data.months_per_borrower,
            "derived": data.account_month_rows,
            "difference": data.account_month_rows
            - data.borrower_rows * data.months_per_borrower,
            "status": "PASS",
            "claim_id": "W1-CLM-002",
            "evidence": "Gate 2 wide_to_long_rows",
        },
        {
            "metric": "Wide bills → long bills",
            "source": data.bill_source_total,
            "derived": data.bill_derived_total,
            "difference": data.bill_difference,
            "status": "PASS",
            "claim_id": "W1-CLM-003",
            "evidence": "Gate 2 bill_amounts",
        },
        {
            "metric": "Wide payments → long payments",
            "source": data.payment_source_total,
            "derived": data.payment_derived_total,
            "difference": data.payment_difference,
            "status": "PASS",
            "claim_id": "W1-CLM-003",
            "evidence": "Gate 2 payment_amounts",
        },
    )


def _format_sheet(
    sheet: Any,
    *,
    freeze: str,
    widths: Sequence[int],
) -> None:
    sheet.freeze_panes = freeze
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[_excel_column(index)].width = width
    sheet.sheet_view.showGridLines = False


def _add_table(sheet: Any, name: str) -> None:
    table = Table(displayName=name, ref=sheet.dimensions)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)


def _excel_column(index: int) -> str:
    result = ""
    while index:
        index, remainder = divmod(index - 1, 26)
        result = chr(65 + remainder) + result
    return result


def _blank_slide(presentation: Presentation) -> Any:
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = RGBColor(255, 255, 255)
    return slide


def _add_banner(slide: Any) -> None:
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0),
        Inches(0),
        Inches(13.333),
        Inches(0.38),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = _ppt_rgb(AMBER)
    shape.line.fill.background()
    frame = shape.text_frame
    frame.clear()
    paragraph = frame.paragraphs[0]
    paragraph.text = INTERNAL_STATUS
    paragraph.alignment = PP_ALIGN.CENTER
    paragraph.runs[0].font.size = Pt(10)
    paragraph.runs[0].font.bold = True
    paragraph.runs[0].font.color.rgb = _ppt_rgb(RED)


def _add_title(slide: Any, text: str) -> None:
    _add_text(slide, 0.75, 0.65, 11.9, 0.65, text, 24, bold=True, color=NAVY)


def _add_text(
    slide: Any,
    x: float,
    y: float,
    width: float,
    height: float,
    text: str,
    size: int,
    *,
    bold: bool = False,
    color: str = DARK,
) -> Any:
    box = slide.shapes.add_textbox(
        Inches(x),
        Inches(y),
        Inches(width),
        Inches(height),
    )
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = _ppt_rgb(color)
    return box


def _add_card(
    slide: Any,
    x: float,
    y: float,
    width: float,
    height: float,
    label: str,
    value: str,
) -> None:
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(width),
        Inches(height),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = _ppt_rgb(LIGHT_BLUE)
    shape.line.color.rgb = _ppt_rgb(BLUE)
    frame = shape.text_frame
    frame.clear()
    frame.margin_left = Inches(0.15)
    frame.margin_right = Inches(0.15)
    p1 = frame.paragraphs[0]
    p1.text = value
    p1.alignment = PP_ALIGN.CENTER
    p1.runs[0].font.size = Pt(25)
    p1.runs[0].font.bold = True
    p1.runs[0].font.color.rgb = _ppt_rgb(NAVY)
    p2 = frame.add_paragraph()
    p2.text = label
    p2.alignment = PP_ALIGN.CENTER
    p2.runs[0].font.size = Pt(11)
    p2.runs[0].font.color.rgb = _ppt_rgb(DARK)


def _add_ppt_table(
    slide: Any,
    rows: Sequence[Sequence[str]],
    x: float,
    y: float,
    width: float,
    height: float,
    column_widths: Sequence[float],
) -> None:
    shape = slide.shapes.add_table(
        len(rows),
        len(rows[0]),
        Inches(x),
        Inches(y),
        Inches(width),
        Inches(height),
    )
    table = shape.table
    for index, column_width in enumerate(column_widths):
        table.columns[index].width = Inches(column_width)
    for row_index, row in enumerate(rows):
        for column_index, value in enumerate(row):
            cell = table.cell(row_index, column_index)
            cell.text = str(value)
            cell.margin_left = Inches(0.06)
            cell.margin_right = Inches(0.06)
            cell.margin_top = Inches(0.03)
            cell.margin_bottom = Inches(0.03)
            cell.fill.solid()
            cell.fill.fore_color.rgb = _ppt_rgb(
                NAVY if row_index == 0 else (PALE if row_index % 2 == 0 else WHITE)
            )
            for paragraph in cell.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(10)
                    run.font.bold = row_index == 0
                    run.font.color.rgb = _ppt_rgb(WHITE if row_index == 0 else DARK)


def _add_footer(slide: Any, data: OfficeData, page: int) -> None:
    _add_text(
        slide,
        0.55,
        7.12,
        12.2,
        0.22,
        (
            f"{data.run_id} | definition {data.definition_version} | "
            f"{INTERNAL_STATUS} | {page}"
        ),
        8,
        color="6B7785",
    )


def _add_docx_table(
    document: Document,
    rows: Sequence[Sequence[str]],
    widths: Sequence[float],
) -> Any:
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    table.autofit = False
    for row_index, row in enumerate(rows):
        for column_index, value in enumerate(row):
            cell = table.cell(row_index, column_index)
            cell.text = str(value)
            cell.width = DocxInches(widths[column_index])
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = DocxPt(8)
                    run.bold = row_index == 0
                    if row_index == 0:
                        run.font.color.rgb = _docx_rgb(WHITE)
            if row_index == 0:
                shading = cell._tc.get_or_add_tcPr()
                fill = shading.makeelement(
                    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}shd",
                    {
                        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fill": NAVY
                    },
                )
                shading.append(fill)
    document.add_paragraph()
    return table


def _excel_rows(sheet: Any, headers: Sequence[str]) -> list[dict[str, Any]]:
    actual_headers = [cell.value for cell in sheet[1]]
    indexes: dict[str, int] = {}
    for header in headers:
        if header not in actual_headers:
            raise OfficePackError(f"Excel sheet {sheet.title} lacks {header}")
        indexes[header] = actual_headers.index(header) + 1
    return [
        {
            header: sheet.cell(row_index, column_index).value
            for header, column_index in indexes.items()
        }
        for row_index in range(2, sheet.max_row + 1)
    ]


def _ppt_table(
    tables: Sequence[Sequence[Sequence[str]]],
    first_header: str,
) -> Sequence[Sequence[str]]:
    for table in tables:
        if table and table[0] and table[0][0] == first_header:
            return table
    raise OfficePackError(f"PowerPoint table {first_header!r} is missing")


def _docx_table(document: Document, first_header: str) -> Any:
    for table in document.tables:
        if table.rows and table.rows[0].cells[0].text == first_header:
            return table
    raise OfficePackError(f"Word table {first_header!r} is missing")


def _normalize_ooxml(path: Path, core_timestamp: str) -> None:
    normalized = _normalized_zip_bytes(path.read_bytes(), core_timestamp)
    path.write_bytes(normalized)


def _normalized_zip_bytes(payload: bytes, core_timestamp: str) -> bytes:
    source_buffer = io.BytesIO(payload)
    target_buffer = io.BytesIO()
    with zipfile.ZipFile(source_buffer, "r") as source:
        with zipfile.ZipFile(target_buffer, "w") as target:
            for name in sorted(source.namelist()):
                source_info = source.getinfo(name)
                content = source.read(name)
                if name == "docProps/core.xml":
                    content = _normalize_core_timestamps(
                        content,
                        core_timestamp,
                    )
                if name.lower().endswith((".xlsx", ".xlsm")):
                    try:
                        content = _normalized_zip_bytes(content, core_timestamp)
                    except zipfile.BadZipFile:
                        pass
                info = zipfile.ZipInfo(name, FIXED_ZIP_TIMESTAMP)
                info.compress_type = source_info.compress_type
                info.external_attr = source_info.external_attr
                info.internal_attr = source_info.internal_attr
                info.create_system = 0
                info.flag_bits = source_info.flag_bits
                target.writestr(info, content)
    return target_buffer.getvalue()


def _normalize_core_timestamps(content: bytes, timestamp: str) -> bytes:
    text = content.decode("utf-8")
    for tag in ("created", "modified"):
        text = re.sub(
            rf"(<dcterms:{tag}\b[^>]*>)[^<]*(</dcterms:{tag}>)",
            rf"\g<1>{timestamp}\g<2>",
            text,
        )
    return text.encode("utf-8")


def _copy_new(source: Path, destination: Path) -> None:
    with source.open("rb") as input_file, destination.open("xb") as output_file:
        shutil.copyfileobj(input_file, output_file)


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for block in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _portable_path(path: Path) -> str:
    try:
        return path.resolve().relative_to(REPOSITORY_ROOT).as_posix()
    except ValueError:
        return path.name


def _safe_token(run_id: str) -> str:
    token = re.sub(r"[^A-Za-z0-9._-]", "_", run_id)
    if not token or token.startswith("."):
        raise OfficePackError("run ID cannot form a safe output filename")
    return token[:128]


def _parse_datetime(value: str) -> datetime:
    try:
        parsed = datetime.fromisoformat(value)
    except ValueError as exc:
        raise OfficePackError("validation generated_at must be ISO-8601") from exc
    if parsed.tzinfo is None:
        raise OfficePackError("validation generated_at must include a timezone")
    return parsed.astimezone(timezone.utc).replace(tzinfo=None)


def _mapping(
    source: Mapping[str, Any],
    field: str,
    path: str,
) -> Mapping[str, Any]:
    value = source.get(field)
    if not isinstance(value, Mapping):
        raise OfficePackError(f"{path}.{field} must be an object")
    return value


def _text(source: Mapping[str, Any], field: str, path: str) -> str:
    value = source.get(field)
    if not isinstance(value, str) or not value.strip():
        raise OfficePackError(f"{path}.{field} must be a non-blank string")
    return value


def _positive_int(source: Mapping[str, Any], field: str, path: str) -> int:
    value = _integer_value(source, field, path)
    if value <= 0:
        raise OfficePackError(f"{path}.{field} must be positive")
    return value


def _nonnegative_int(source: Mapping[str, Any], field: str, path: str) -> int:
    value = _integer_value(source, field, path)
    if value < 0:
        raise OfficePackError(f"{path}.{field} must be non-negative")
    return value


def _integer_value(source: Mapping[str, Any], field: str, path: str) -> int:
    value = source.get(field)
    if not isinstance(value, int) or isinstance(value, bool):
        raise OfficePackError(f"{path}.{field} must be an integer")
    return value


def _decimal_text(value: object, path: str) -> str:
    if isinstance(value, bool) or not isinstance(value, (int, float, str)):
        raise OfficePackError(f"{path} must be a finite decimal")
    try:
        decimal = Decimal(str(value))
    except Exception as exc:
        raise OfficePackError(f"{path} must be a finite decimal") from exc
    if not decimal.is_finite():
        raise OfficePackError(f"{path} must be a finite decimal")
    return str(decimal)


def _equal(left: object, right: object, label: str) -> None:
    if left != right:
        raise OfficePackError(f"{label} mismatch: {left!r} != {right!r}")


def _integer(value: int) -> str:
    return f"{value:,}"


def _percent(value: object) -> str:
    rate = Decimal(str(value))
    percent = (rate * Decimal("100")).quantize(
        Decimal("0.01"),
        rounding=ROUND_HALF_UP,
    )
    return f"{percent}%"


def _ppt_rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def _docx_rgb(value: str) -> Any:
    from docx.shared import RGBColor as DocxRGBColor

    return DocxRGBColor.from_string(value)


def parse_args(argv: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--validation", type=Path, default=DEFAULT_VALIDATION)
    parser.add_argument(
        "--claim-manifest",
        type=Path,
        default=DEFAULT_CLAIM_MANIFEST,
    )
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    parser.add_argument(
        "--check-only",
        action="store_true",
        help="validate inputs without writing Office artifacts",
    )
    return parser.parse_args(argv)


def main(argv: Sequence[str] | None = None) -> int:
    args = parse_args(argv)
    try:
        if args.check_only:
            data = load_office_data(args.validation, args.claim_manifest)
            print(
                json.dumps(
                    {
                        "status": "eligible_internal_office_pack",
                        "run_id": data.run_id,
                        "definition_version": data.definition_version,
                        "signal_bucket_count": len(data.signal_rows),
                        "risk_band_count": len(data.band_rows),
                        "publication_status": "pending_gate_5_and_human_approval",
                    },
                    ensure_ascii=False,
                )
            )
            return 0
        result = generate_office_pack(
            args.validation,
            args.claim_manifest,
            args.output_dir,
        )
    except (OSError, ValueError, json.JSONDecodeError, OfficePackError) as exc:
        print(
            json.dumps(
                {"status": "blocked", "reason": str(exc)},
                ensure_ascii=False,
            )
        )
        return 2
    print(
        json.dumps(
            {
                "status": "generated_verified_internal_office_pack",
                "run_id": result.run_id,
                "excel": str(result.excel_path),
                "powerpoint": str(result.powerpoint_path),
                "word": str(result.word_path),
                "manifest": str(result.manifest_path),
                "publication_status": "pending_gate_5_and_human_approval",
            },
            ensure_ascii=False,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
