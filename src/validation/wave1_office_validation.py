"""Independently validate the Wave 1 XLSX/PPTX/DOCX evidence pack.

This module reopens each OOXML package without importing the Office Builder.
The independently validated Wave 1 JSON and run-specific claim manifest are
the numerical and claim authorities.  A passing Gate 5 packet is consistency
evidence only; human publication approval remains required.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import shutil
from dataclasses import asdict, dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Mapping, Sequence

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

from .gate_evidence import GateEvidenceValidator, load_json_object


INTERNAL_STATUS = "INTERNAL ONLY — PENDING GATE 5 AND HUMAN PUBLICATION APPROVAL"
IFRS9_LIMITATION = (
    "Public-data descriptive prototype; not a production IFRS 9 rating, "
    "Stage, PD, LGD, EAD, or ECL result."
)
LIMITATIONS = (
    IFRS9_LIMITATION,
    "Observed default rates are descriptive associations, not causal effects.",
    (
        "Risk bands are explainable prototype rules, not approved bank "
        "ratings or IFRS 9 Stage assignments."
    ),
    (
        "The independent validation artifact contains the aggregate SQL "
        "test count, not all individual test-result records."
    ),
    (
        "Displayed percentages are rounded to two decimals; exact counts "
        "and exact decimal-rate strings remain the audit authority."
    ),
    (
        "Gate 1–3 machine evidence does not constitute Gate 5 review or "
        "human approval for external publication."
    ),
)
REQUIRED_CLAIM_IDS = tuple(f"W1-CLM-{number:03d}" for number in range(1, 7))
REQUIRED_XLSX_SHEETS = (
    "Control",
    "Reconciliation",
    "DQ_Warnings",
    "Gate_Checks",
    "Signal_Buckets",
    "Risk_Bands",
    "Limitations",
    "Evidence",
)
REQUIRED_DOCX_HEADINGS = (
    "Technical summary",
    "Reconciled populations protect the signal denominator",
    "Data-quality controls passed with one reporting limitation",
    "Observed default rates differ by initial risk band",
    "All validated signal-bucket outcomes remain auditable",
    "Scope, data, and metric definitions",
    "Methodology and reporting controls",
    "Limitations, uncertainty, and robustness checks",
    "Claim-to-evidence mapping",
    "Recommended next steps",
    "Further questions",
)
REQUIRED_PPTX_TITLES = (
    "Credit Risk Signal Lab",
    "Validated Wave 1 evidence is reproducible and reviewable",
    "The relational transformation preserves rows and amounts",
    "Data-quality controls passed; detailed test IDs remain upstream",
    "Observed default rates separate across initial risk bands",
    "More delinquent months align with higher observed default",
    "Traceability is complete; publication approval is not",
)


@dataclass(frozen=True)
class OfficeIssue:
    code: str
    path: str
    message: str


@dataclass(frozen=True)
class OfficeValidationReport:
    schema_version: str
    run_id: str | None
    generated_at: str
    status: str
    mismatch_count: int
    unresolved_claim_count: int
    limitations_documented: bool
    human_publication_approval: str
    visual_render: Mapping[str, Any]
    artifact_checks: Mapping[str, Any]
    issues: tuple[OfficeIssue, ...]
    gate_packet: Mapping[str, Any] | None

    def to_dict(self) -> dict[str, Any]:
        return {
            "schema_version": self.schema_version,
            "run_id": self.run_id,
            "generated_at": self.generated_at,
            "status": self.status,
            "mismatch_count": self.mismatch_count,
            "unresolved_claim_count": self.unresolved_claim_count,
            "limitations_documented": self.limitations_documented,
            "human_publication_approval": self.human_publication_approval,
            "visual_render": dict(self.visual_render),
            "artifact_checks": dict(self.artifact_checks),
            "issues": [asdict(issue) for issue in self.issues],
            "gate_packet": dict(self.gate_packet) if self.gate_packet else None,
        }


def validate_office_pack(
    office_manifest_path: Path,
    validation_path: Path,
    claim_manifest_path: Path,
    evidence_map_path: Path,
    *,
    artifact_root: Path,
    generated_at: str | None = None,
) -> OfficeValidationReport:
    """Reopen and reconcile one immutable Office pack."""

    root = artifact_root.resolve()
    issues: list[OfficeIssue] = []
    artifact_checks: dict[str, Any] = {}
    timestamp = generated_at or datetime.now(timezone.utc).isoformat()
    run_id: str | None = None
    unresolved_claim_count = len(REQUIRED_CLAIM_IDS)
    limitations_documented = False
    gate_packet: Mapping[str, Any] | None = None

    try:
        manifest = load_json_object(office_manifest_path)
        validation = load_json_object(validation_path)
        claims = load_json_object(claim_manifest_path)
        evidence_map = evidence_map_path.read_text(encoding="utf-8")
    except (OSError, ValueError, json.JSONDecodeError) as exc:
        _issue(issues, "input_unreadable", "$", str(exc))
        return _report(
            run_id,
            timestamp,
            issues,
            unresolved_claim_count,
            limitations_documented,
            artifact_checks,
            gate_packet,
        )

    canonical = _canonical_evidence(validation, issues)
    if canonical:
        run_id = canonical["run_id"]
    claim_rows, unresolved_claim_count = _validate_claims(
        claims,
        canonical,
        validation_path,
        evidence_map,
        evidence_map_path,
        issues,
    )
    artifact_paths = _validate_manifest(
        manifest,
        canonical,
        claim_rows,
        office_manifest_path,
        validation_path,
        claim_manifest_path,
        root,
        issues,
    )
    limitations_documented = (
        isinstance(manifest.get("limitations"), list)
        and tuple(manifest["limitations"]) == LIMITATIONS
    )

    excel_receipt: Mapping[str, Any] = {}
    pptx_receipt: Mapping[str, Any] = {}
    docx_receipt: Mapping[str, Any] = {}
    if canonical and claim_rows:
        excel_path = artifact_paths.get("excel")
        pptx_path = artifact_paths.get("powerpoint")
        docx_path = artifact_paths.get("word")
        if excel_path is not None:
            excel_receipt = _validate_excel(
                excel_path,
                canonical,
                claim_rows,
                validation_path,
                claim_manifest_path,
                root,
                issues,
            )
            artifact_checks["excel"] = excel_receipt
        if pptx_path is not None:
            pptx_receipt = _validate_powerpoint(pptx_path, canonical, issues)
            artifact_checks["powerpoint"] = pptx_receipt
        if docx_path is not None:
            docx_receipt = _validate_word(docx_path, canonical, claim_rows, issues)
            artifact_checks["word"] = docx_receipt
        _validate_cross_artifact(
            excel_receipt,
            pptx_receipt,
            docx_receipt,
            canonical,
            issues,
        )

    mismatch_count = len(issues)
    if mismatch_count == 0 and unresolved_claim_count == 0 and run_id is not None:
        gate_packet = _gate_5_packet(
            run_id,
            timestamp,
            office_manifest_path=office_manifest_path,
            evidence_map_path=evidence_map_path,
            claim_count=len(claim_rows),
            artifact_count=3,
            limitation_count=len(LIMITATIONS),
        )
        contract = GateEvidenceValidator(artifact_root=root).validate(
            gate_packet,
            expected_gate=5,
            expected_run_id=run_id,
        )
        if not contract.valid:
            for item in contract.issues:
                _issue(
                    issues,
                    f"gate_contract_{item.code}",
                    item.path,
                    item.message,
                )
            gate_packet = None

    return _report(
        run_id,
        timestamp,
        issues,
        unresolved_claim_count,
        limitations_documented,
        artifact_checks,
        gate_packet,
    )


def _canonical_evidence(
    validation: Mapping[str, Any],
    issues: list[OfficeIssue],
) -> dict[str, Any]:
    if validation.get("schema_version") != "1.0" or validation.get("status") != "pass":
        _issue(
            issues,
            "validation_not_pass",
            "$.validation",
            "independent Wave 1 validation must be schema 1.0 with status pass",
        )
        return {}
    if validation.get("issues") != []:
        _issue(
            issues,
            "validation_has_issues",
            "$.validation.issues",
            "independent Wave 1 validation issues must be empty",
        )
    run_id = validation.get("run_id")
    generated_at = validation.get("generated_at")
    if not isinstance(run_id, str) or not run_id:
        _issue(issues, "invalid_run_id", "$.validation.run_id", "run_id is required")
        return {}
    packets_value = validation.get("gate_packets")
    contracts = validation.get("contract_reports")
    findings = validation.get("verified_findings")
    if not isinstance(packets_value, list) or not isinstance(findings, Mapping):
        _issue(
            issues,
            "invalid_validation_shape",
            "$.validation",
            "gate_packets and verified_findings are required",
        )
        return {}
    packets = {
        packet.get("gate"): packet
        for packet in packets_value
        if isinstance(packet, Mapping)
    }
    if set(packets) != {1, 2, 3}:
        _issue(
            issues,
            "gate_packet_coverage",
            "$.validation.gate_packets",
            "exactly Gates 1-3 are required",
        )
        return {}
    if not isinstance(contracts, list) or any(
        not isinstance(item, Mapping)
        or item.get("valid") is not True
        or item.get("run_id") != run_id
        for item in contracts
    ):
        _issue(
            issues,
            "contract_report_mismatch",
            "$.validation.contract_reports",
            "all independent contract reports must pass for the same run",
        )
    for gate, packet in packets.items():
        if packet.get("run_id") != run_id or packet.get("status") != "pass":
            _issue(
                issues,
                "gate_run_status_mismatch",
                f"$.validation.gate_packets[{gate}]",
                "each Gate 1-3 packet must pass for the shared run",
            )
    checks = {gate: _checks(packet, gate, issues) for gate, packet in packets.items()}
    if any(not value for value in checks.values()):
        return {}

    g1, g2, g3 = checks[1], checks[2], checks[3]
    source = _evidence(g1, "source_registered", issues)
    sha = _evidence(g1, "sha256_recorded", issues)
    rows = _evidence(g2, "wide_to_long_rows", issues)
    bills = _evidence(g2, "bill_amounts", issues)
    payments = _evidence(g2, "payment_amounts", issues)
    version = _evidence(g3, "definition_version", issues)
    signal_count = _evidence(g3, "signal_count", issues)
    signal_rows = findings.get("signal_outcomes")
    band_rows = findings.get("risk_band_outcomes")
    if not isinstance(signal_rows, list) or not isinstance(band_rows, list):
        _issue(
            issues,
            "outcome_rows_missing",
            "$.validation.verified_findings",
            "signal and risk-band outcome arrays are required",
        )
        return {}

    canonical = {
        "run_id": run_id,
        "generated_at": generated_at,
        "definition_version": version.get("version"),
        "source_id": source.get("source_id"),
        "source_sha256": sha.get("sha256"),
        "source_rows": findings.get("source_rows"),
        "raw_rows": findings.get("raw_rows"),
        "borrower_rows": findings.get("borrower_rows"),
        "account_month_rows": findings.get("account_month_rows"),
        "months_per_borrower": rows.get("months_per_borrower"),
        "bill_source_total": bills.get("source_total"),
        "bill_derived_total": bills.get("derived_total"),
        "bill_difference": bills.get("difference"),
        "payment_source_total": payments.get("source_total"),
        "payment_derived_total": payments.get("derived_total"),
        "payment_difference": payments.get("difference"),
        "sql_test_count": findings.get("sql_test_count"),
        "sql_failed_count": findings.get("sql_failed_count"),
        "signal_count": signal_count.get("count"),
        "signal_rows": signal_rows,
        "band_rows": band_rows,
        "gate_packets": packets,
    }
    expected_counts = {
        "source_rows": 30_000,
        "raw_rows": 30_000,
        "borrower_rows": 30_000,
        "account_month_rows": 180_000,
        "months_per_borrower": 6,
        "sql_test_count": 15,
        "sql_failed_count": 0,
        "signal_count": 7,
    }
    for field, expected in expected_counts.items():
        if canonical[field] != expected:
            _issue(
                issues,
                "canonical_value_mismatch",
                f"$.validation.{field}",
                f"{field} must equal {expected!r}",
            )
    if len(signal_rows) != 23 or len(band_rows) != 3:
        _issue(
            issues,
            "canonical_outcome_count",
            "$.validation.verified_findings",
            "canonical evidence requires 23 signal rows and 3 risk-band rows",
        )
    for row in [*signal_rows, *band_rows]:
        if not isinstance(row, Mapping):
            _issue(
                issues,
                "canonical_outcome_shape",
                "$.validation.verified_findings",
                "every outcome row must be an object",
            )
            continue
        count = row.get("sample_count")
        defaults = row.get("default_count")
        rate = row.get("observed_default_rate")
        try:
            calculated = str(int(defaults) / int(count))
            valid_rate = abs(float(rate) - float(calculated)) <= 1e-15
        except (TypeError, ValueError, ZeroDivisionError):
            valid_rate = False
        if not valid_rate:
            _issue(
                issues,
                "canonical_rate_mismatch",
                "$.validation.verified_findings",
                "observed rate does not reconcile to default_count / sample_count",
            )
    return canonical


def _validate_claims(
    manifest: Mapping[str, Any],
    canonical: Mapping[str, Any],
    validation_path: Path,
    evidence_map: str,
    evidence_map_path: Path,
    issues: list[OfficeIssue],
) -> tuple[list[Mapping[str, Any]], int]:
    value = manifest.get("claims")
    if not canonical or not isinstance(value, list):
        _issue(
            issues,
            "claim_manifest_shape",
            "$.claim_manifest.claims",
            "claim manifest must contain a claims list",
        )
        return [], len(REQUIRED_CLAIM_IDS)
    by_id: dict[str, Mapping[str, Any]] = {}
    validation_hash = _sha256(validation_path)
    for index, row in enumerate(value):
        if not isinstance(row, Mapping):
            _issue(
                issues,
                "claim_not_object",
                f"$.claim_manifest.claims[{index}]",
                "claim must be an object",
            )
            continue
        claim_id = row.get("claim_id")
        if not isinstance(claim_id, str) or claim_id in by_id:
            _issue(
                issues,
                "claim_id_invalid",
                f"$.claim_manifest.claims[{index}].claim_id",
                "claim ID must be unique",
            )
            continue
        by_id[claim_id] = row
        if (
            row.get("run_id") != canonical["run_id"]
            or row.get("required_gate_status") != "pass"
            or row.get("validation_artifact_sha256") != validation_hash
            or not isinstance(row.get("json_pointers"), list)
            or not row["json_pointers"]
        ):
            _issue(
                issues,
                "claim_evidence_mismatch",
                f"$.claim_manifest.claims[{index}]",
                "claim run, gate, validation hash, or evidence pointers do not reconcile",
            )
    if set(by_id) != set(REQUIRED_CLAIM_IDS):
        _issue(
            issues,
            "claim_coverage_mismatch",
            "$.claim_manifest.claims",
            "claim manifest must contain exactly W1-CLM-001 through W1-CLM-006",
        )
    unresolved = 0
    for claim_id in REQUIRED_CLAIM_IDS:
        pattern = re.compile(rf"^\|\s*{re.escape(claim_id)}\s*\|", re.MULTILINE)
        if claim_id not in by_id or pattern.search(evidence_map) is None:
            unresolved += 1
    if unresolved:
        _issue(
            issues,
            "evidence_map_unresolved_claims",
            str(evidence_map_path),
            f"{unresolved} required claim IDs are absent from the evidence map",
        )
    publication = manifest.get("publication")
    if not isinstance(publication, Mapping) or (
        publication.get("status") != "pending"
        or publication.get("required_gate") != 5
        or publication.get("human_approval_required") is not True
    ):
        _issue(
            issues,
            "claim_publication_control",
            "$.claim_manifest.publication",
            "claim publication status must remain pending human approval",
        )
    return [by_id[item] for item in REQUIRED_CLAIM_IDS if item in by_id], unresolved


def _validate_manifest(
    manifest: Mapping[str, Any],
    canonical: Mapping[str, Any],
    claim_rows: Sequence[Mapping[str, Any]],
    office_manifest_path: Path,
    validation_path: Path,
    claim_manifest_path: Path,
    root: Path,
    issues: list[OfficeIssue],
) -> dict[str, Path]:
    paths: dict[str, Path] = {}
    if not canonical:
        return paths
    expected_scalars = {
        "schema_version": "1.0",
        "pack_type": "wave1_internal_office_pack",
        "run_id": canonical["run_id"],
        "signal_definition_version": canonical["definition_version"],
        "validation_generated_at": canonical["generated_at"],
    }
    for field, expected in expected_scalars.items():
        if manifest.get(field) != expected:
            _issue(
                issues,
                "manifest_identity_mismatch",
                f"$.office_manifest.{field}",
                f"{field} does not match canonical evidence",
            )
    publication = manifest.get("publication")
    if not isinstance(publication, Mapping) or publication != {
        "status": "pending",
        "required_gate": 5,
        "human_approval_required": True,
        "label": INTERNAL_STATUS,
    }:
        _issue(
            issues,
            "manifest_publication_control",
            "$.office_manifest.publication",
            "Office publication must remain pending Gate 5 and human approval",
        )
    if manifest.get("source") != {
        "source_id": canonical["source_id"],
        "sha256": canonical["source_sha256"],
    }:
        _issue(
            issues,
            "manifest_source_mismatch",
            "$.office_manifest.source",
            "Office source identity does not match canonical evidence",
        )
    expected_cross = {
        "source_rows": canonical["source_rows"],
        "raw_rows": canonical["raw_rows"],
        "borrower_rows": canonical["borrower_rows"],
        "account_month_rows": canonical["account_month_rows"],
        "months_per_borrower": canonical["months_per_borrower"],
        "bill_difference": canonical["bill_difference"],
        "payment_difference": canonical["payment_difference"],
        "sql_test_count": canonical["sql_test_count"],
        "sql_failed_count": canonical["sql_failed_count"],
        "signal_count": canonical["signal_count"],
        "signal_bucket_count": len(canonical["signal_rows"]),
        "risk_band_count": len(canonical["band_rows"]),
        "signal_rows": canonical["signal_rows"],
        "risk_band_rows": canonical["band_rows"],
    }
    if manifest.get("cross_artifact_values") != expected_cross:
        _issue(
            issues,
            "manifest_values_mismatch",
            "$.office_manifest.cross_artifact_values",
            "manifest cross-artifact values do not exactly match canonical evidence",
        )
    if manifest.get("claims") != list(REQUIRED_CLAIM_IDS):
        _issue(
            issues,
            "manifest_claims_mismatch",
            "$.office_manifest.claims",
            "Office manifest must list all six claim IDs in order",
        )
    if manifest.get("limitations") != list(LIMITATIONS):
        _issue(
            issues,
            "manifest_limitations_mismatch",
            "$.office_manifest.limitations",
            "Office manifest limitations do not match approved limitations",
        )
    expected_inputs = {
        "independent_validation": (validation_path, _sha256(validation_path)),
        "claim_manifest": (claim_manifest_path, _sha256(claim_manifest_path)),
    }
    inputs = manifest.get("inputs")
    if not isinstance(inputs, list) or len(inputs) != 2:
        _issue(
            issues,
            "manifest_inputs_shape",
            "$.office_manifest.inputs",
            "manifest must contain exactly two input records",
        )
    else:
        for index, row in enumerate(inputs):
            if not isinstance(row, Mapping) or row.get("input_type") not in expected_inputs:
                _issue(
                    issues,
                    "manifest_input_invalid",
                    f"$.office_manifest.inputs[{index}]",
                    "manifest input type is invalid",
                )
                continue
            expected_path, expected_hash = expected_inputs[str(row["input_type"])]
            if (
                _resolve_portable(row.get("path"), root, issues) != expected_path.resolve()
                or row.get("sha256") != expected_hash
            ):
                _issue(
                    issues,
                    "manifest_input_mismatch",
                    f"$.office_manifest.inputs[{index}]",
                    "manifest input path/hash does not match current canonical input",
                )
    artifacts = manifest.get("artifacts")
    if not isinstance(artifacts, list) or len(artifacts) != 3:
        _issue(
            issues,
            "manifest_artifacts_shape",
            "$.office_manifest.artifacts",
            "manifest must contain exactly three Office artifacts",
        )
        return paths
    for index, row in enumerate(artifacts):
        path = f"$.office_manifest.artifacts[{index}]"
        if not isinstance(row, Mapping):
            _issue(issues, "artifact_record_invalid", path, "artifact record must be object")
            continue
        artifact_type = row.get("artifact_type")
        if artifact_type not in {"excel", "powerpoint", "word"} or artifact_type in paths:
            _issue(
                issues,
                "artifact_type_invalid",
                f"{path}.artifact_type",
                "artifact type must be unique Excel, PowerPoint, or Word",
            )
            continue
        resolved = _resolve_portable(row.get("path"), root, issues)
        if resolved is None or not resolved.is_file():
            _issue(issues, "artifact_missing", f"{path}.path", "artifact file is missing")
            continue
        paths[str(artifact_type)] = resolved
        if (
            row.get("sha256") != _sha256(resolved)
            or row.get("bytes") != resolved.stat().st_size
            or row.get("reopened") is not True
            or row.get("cross_artifact_status") != "pass"
        ):
            _issue(
                issues,
                "artifact_hash_or_receipt_mismatch",
                path,
                "artifact hash, size, reopened flag, or receipt status does not match",
            )
    if set(paths) != {"excel", "powerpoint", "word"}:
        _issue(
            issues,
            "artifact_coverage_mismatch",
            str(office_manifest_path),
            "all three artifact paths must resolve inside the repository",
        )
    return paths


def _validate_excel(
    path: Path,
    canonical: Mapping[str, Any],
    claims: Sequence[Mapping[str, Any]],
    validation_path: Path,
    claim_manifest_path: Path,
    root: Path,
    issues: list[OfficeIssue],
) -> dict[str, Any]:
    receipt: dict[str, Any] = {}
    try:
        workbook = load_workbook(path, read_only=False, data_only=False)
    except Exception as exc:  # OOXML parser exceptions vary by dependency version.
        _issue(issues, "xlsx_unreadable", str(path), str(exc))
        return receipt
    if tuple(workbook.sheetnames) != REQUIRED_XLSX_SHEETS:
        _issue(
            issues,
            "xlsx_sheet_mismatch",
            str(path),
            "XLSX sheet names/order do not match the required contract",
        )
    try:
        control = {
            str(row[0]): row[1]
            for row in workbook["Control"].iter_rows(min_row=4, values_only=True)
            if row[0] is not None
        }
        expected_control = {
            "Run ID": canonical["run_id"],
            "Signal definition version": canonical["definition_version"],
            "Validation generated at": canonical["generated_at"],
            "Source ID": canonical["source_id"],
            "Source SHA-256": canonical["source_sha256"],
            "Validation artifact": _portable_to_root(validation_path, root),
            "Validation SHA-256": _sha256(validation_path),
            "Claim manifest": _portable_to_root(claim_manifest_path, root),
            "Claim manifest SHA-256": _sha256(claim_manifest_path),
            "Publication": "pending Gate 5 and human approval",
            "Limitation": IFRS9_LIMITATION,
        }
        _match(control, expected_control, issues, "xlsx_control", str(path))

        reconciliation = _sheet_dicts(workbook["Reconciliation"])
        expected_reconciliation = _reconciliation_rows(canonical)
        _match(
            [_stringify(row) for row in reconciliation],
            [_stringify(row) for row in expected_reconciliation],
            issues,
            "xlsx_reconciliation",
            str(path),
        )
        signal_rows = _sheet_dicts(workbook["Signal_Buckets"])
        expected_signals = _excel_signal_rows(canonical)
        _match_excel_outcomes(
            signal_rows,
            expected_signals,
            issues,
            "xlsx_signal_values",
            str(path),
        )
        band_rows = _sheet_dicts(workbook["Risk_Bands"])
        expected_bands = _excel_band_rows(canonical)
        _match_excel_outcomes(
            band_rows,
            expected_bands,
            issues,
            "xlsx_band_values",
            str(path),
        )

        gate_rows = _sheet_dicts(workbook["Gate_Checks"])
        expected_gates = _excel_gate_rows(canonical)
        _match(gate_rows, expected_gates, issues, "xlsx_gate_values", str(path))
        dq_rows = _sheet_dicts(workbook["DQ_Warnings"])
        _match(
            [_stringify(row) for row in dq_rows],
            [_stringify(row) for row in _dq_rows(canonical)],
            issues,
            "xlsx_dq_values",
            str(path),
        )
        limitation_rows = _sheet_dicts(workbook["Limitations"])
        expected_limitations = [
            {"Type": "Limitation", "Detail": item} for item in LIMITATIONS
        ] + [{"Type": "Publication control", "Detail": INTERNAL_STATUS}]
        _match(
            limitation_rows,
            expected_limitations,
            issues,
            "xlsx_limitations",
            str(path),
        )
        evidence_rows = _sheet_dicts(workbook["Evidence"])
        _match(
            evidence_rows,
            _excel_claim_rows(claims, canonical, validation_path),
            issues,
            "xlsx_claim_values",
            str(path),
        )
        if sum(len(sheet._charts) for sheet in workbook.worksheets) != 1:
            _issue(
                issues,
                "xlsx_chart_count",
                str(path),
                "XLSX must contain exactly one structural chart",
            )
        receipt = {
            "run_id": control.get("Run ID"),
            "definition_version": control.get("Signal definition version"),
            "source_sha256": control.get("Source SHA-256"),
            "reconciliation_rows": [_stringify(row) for row in reconciliation],
            "signal_rows": signal_rows,
            "band_rows": band_rows,
            "claim_ids": [row.get("Claim ID") for row in evidence_rows],
            "limitations": [row.get("Detail") for row in limitation_rows[:-1]],
            "publication_label": limitation_rows[-1].get("Detail")
            if limitation_rows
            else None,
        }
    except (KeyError, TypeError, ValueError, IndexError) as exc:
        _issue(issues, "xlsx_contract_error", str(path), str(exc))
    finally:
        workbook.close()
    return receipt


def _validate_powerpoint(
    path: Path,
    canonical: Mapping[str, Any],
    issues: list[OfficeIssue],
) -> dict[str, Any]:
    receipt: dict[str, Any] = {}
    try:
        presentation = Presentation(path)
    except Exception as exc:
        _issue(issues, "pptx_unreadable", str(path), str(exc))
        return receipt
    if len(presentation.slides) != len(REQUIRED_PPTX_TITLES):
        _issue(
            issues,
            "pptx_slide_count",
            str(path),
            f"PowerPoint must contain {len(REQUIRED_PPTX_TITLES)} slides",
        )
    texts_by_slide: list[str] = []
    tables: list[list[list[str]]] = []
    chart_count = 0
    for slide in presentation.slides:
        texts: list[str] = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                texts.append(shape.text)
            if getattr(shape, "has_table", False):
                tables.append(
                    [[cell.text for cell in row.cells] for row in shape.table.rows]
                )
            if getattr(shape, "has_chart", False):
                chart_count += 1
        texts_by_slide.append("\n".join(texts))
    joined = "\n".join(texts_by_slide)
    for index, title in enumerate(REQUIRED_PPTX_TITLES):
        if index >= len(texts_by_slide) or title not in texts_by_slide[index]:
            _issue(
                issues,
                "pptx_title_missing",
                f"{path}#slide-{index + 1}",
                f"required slide title is missing: {title}",
            )
    for value in (
        canonical["run_id"],
        canonical["definition_version"],
        canonical["source_sha256"],
        INTERNAL_STATUS,
        IFRS9_LIMITATION,
        "30,000",
        "180,000",
        str(canonical["sql_test_count"]),
    ):
        if str(value) not in joined:
            _issue(
                issues,
                "pptx_required_value_missing",
                str(path),
                f"PowerPoint is missing required value {value!r}",
            )
    if joined.count(INTERNAL_STATUS) < len(presentation.slides):
        _issue(
            issues,
            "pptx_publication_label",
            str(path),
            "every slide must carry the internal-only publication label",
        )
    if chart_count != 2:
        _issue(
            issues,
            "pptx_chart_count",
            str(path),
            "PowerPoint must contain exactly two structural charts",
        )
    reconciliation = _ppt_table(tables, "Metric")
    expected_reconciliation = [
        [
            str(row["Metric"]),
            str(row["Source / expected"]),
            str(row["Derived / actual"]),
            str(row["Difference"]),
        ]
        for row in _reconciliation_rows(canonical)
    ]
    _match(
        reconciliation[1:] if reconciliation else [],
        expected_reconciliation,
        issues,
        "pptx_reconciliation",
        str(path),
    )
    bands = _ppt_table(tables, "Band")
    expected_bands = [
        [
            str(row["risk_band"]),
            _integer(row["sample_count"]),
            _integer(row["default_count"]),
            _percent(row["observed_default_rate"]),
        ]
        for row in canonical["band_rows"]
    ]
    _match(
        bands[1:] if bands else [],
        expected_bands,
        issues,
        "pptx_band_values",
        str(path),
    )
    selected = [
        row
        for row in canonical["signal_rows"]
        if row["signal_id"] == "delinquent_months_6m"
    ]
    signal = _ppt_table(tables, "6m delinquent months")
    expected_signal = [
        [
            str(row["bucket"]),
            _integer(row["sample_count"]),
            _integer(row["default_count"]),
            _percent(row["observed_default_rate"]),
        ]
        for row in selected
    ]
    _match(
        signal[1:] if signal else [],
        expected_signal,
        issues,
        "pptx_signal_values",
        str(path),
    )
    receipt = {
        "run_id": canonical["run_id"] if canonical["run_id"] in joined else None,
        "definition_version": (
            canonical["definition_version"]
            if canonical["definition_version"] in joined
            else None
        ),
        "source_sha256": (
            canonical["source_sha256"] if canonical["source_sha256"] in joined else None
        ),
        "reconciliation_rows": reconciliation[1:] if reconciliation else [],
        "band_rows": bands[1:] if bands else [],
        "selected_signal_rows": signal[1:] if signal else [],
        "publication_label": INTERNAL_STATUS if INTERNAL_STATUS in joined else None,
        "limitations": [IFRS9_LIMITATION] if IFRS9_LIMITATION in joined else [],
        "slide_count": len(presentation.slides),
        "chart_count": chart_count,
    }
    return receipt


def _validate_word(
    path: Path,
    canonical: Mapping[str, Any],
    claims: Sequence[Mapping[str, Any]],
    issues: list[OfficeIssue],
) -> dict[str, Any]:
    receipt: dict[str, Any] = {}
    try:
        document = Document(path)
    except Exception as exc:
        _issue(issues, "docx_unreadable", str(path), str(exc))
        return receipt
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    joined = "\n".join(paragraphs)
    headings = [
        paragraph.text
        for paragraph in document.paragraphs
        if paragraph.style.name.startswith("Heading")
    ]
    if tuple(headings) != REQUIRED_DOCX_HEADINGS:
        _issue(
            issues,
            "docx_heading_mismatch",
            str(path),
            "Word required headings/order do not match",
        )
    if len(document.tables) != 7:
        _issue(
            issues,
            "docx_table_count",
            str(path),
            "Word must contain exactly seven required tables",
        )
    for value in (
        canonical["run_id"],
        INTERNAL_STATUS,
        IFRS9_LIMITATION,
        "30,000",
        "180,000",
        str(canonical["sql_test_count"]),
    ):
        if str(value) not in joined:
            _issue(
                issues,
                "docx_required_value_missing",
                str(path),
                f"Word is missing required value {value!r}",
            )
    for limitation in LIMITATIONS:
        if limitation not in joined:
            _issue(
                issues,
                "docx_limitation_missing",
                str(path),
                f"Word is missing approved limitation: {limitation}",
            )
    tables = {
        table.rows[0].cells[0].text: table
        for table in document.tables
        if table.rows and table.rows[0].cells
    }
    expected_headers = {
        "Field",
        "Metric",
        "Control",
        "Risk band",
        "Signal ID",
        "Definition",
        "Claim ID",
    }
    if set(tables) != expected_headers:
        _issue(
            issues,
            "docx_table_headers",
            str(path),
            "Word required table headers do not match",
        )
        return receipt
    metadata = _doc_table_dicts(tables["Field"])
    metadata_map = {row["Field"]: row["Value"] for row in metadata}
    expected_metadata = {
        "Run ID": canonical["run_id"],
        "Signal definition version": canonical["definition_version"],
        "Validation generated at": canonical["generated_at"],
        "Source ID": canonical["source_id"],
        "Source SHA-256": canonical["source_sha256"],
        "Publication": "Pending Gate 5 and human approval",
    }
    _match(metadata_map, expected_metadata, issues, "docx_metadata", str(path))

    reconciliation = _doc_table_dicts(tables["Metric"])
    _match(
        [_stringify(row) for row in reconciliation],
        [
            _stringify(
                {
                    key: value
                    for key, value in row.items()
                    if key
                    in {
                        "Metric",
                        "Source / expected",
                        "Derived / actual",
                        "Difference",
                        "Status",
                    }
                }
            )
            for row in _reconciliation_rows(canonical)
        ],
        issues,
        "docx_reconciliation",
        str(path),
    )
    dq = _doc_table_dicts(tables["Control"])
    expected_dq = [
        {
            "Control": str(row["Control ID"]),
            "Status": str(row["Status"]),
            "Observed": str(row["Observed"]),
            "Expectation": str(row["Threshold / expectation"]),
            "Evidence / warning": str(row["Evidence / warning"]),
        }
        for row in _dq_rows(canonical)
    ]
    _match(dq, expected_dq, issues, "docx_dq", str(path))

    bands = _doc_table_dicts(tables["Risk band"])
    expected_bands = [
        {
            "Risk band": str(row["risk_band"]),
            "Sample": str(row["sample_count"]),
            "Defaults": str(row["default_count"]),
            "Exact rate": str(row["observed_default_rate"]),
            "Display rate": _percent(row["observed_default_rate"]),
        }
        for row in canonical["band_rows"]
    ]
    _match(bands, expected_bands, issues, "docx_band_values", str(path))
    signals = _doc_table_dicts(tables["Signal ID"])
    expected_signals = [
        {
            "Signal ID": str(row["signal_id"]),
            "Bucket": str(row["bucket"]),
            "Sample": str(row["sample_count"]),
            "Defaults": str(row["default_count"]),
            "Exact rate": str(row["observed_default_rate"]),
            "Display rate": _percent(row["observed_default_rate"]),
        }
        for row in canonical["signal_rows"]
    ]
    _match(signals, expected_signals, issues, "docx_signal_values", str(path))
    claim_table = _doc_table_dicts(tables["Claim ID"])
    expected_claims = [
        {
            "Claim ID": row["claim_id"],
            "Claim": row["claim"],
            "JSON evidence pointers": " | ".join(row["json_pointers"]),
        }
        for row in claims
    ]
    _match(claim_table, expected_claims, issues, "docx_claim_values", str(path))
    receipt = {
        "run_id": metadata_map.get("Run ID"),
        "definition_version": metadata_map.get("Signal definition version"),
        "source_sha256": metadata_map.get("Source SHA-256"),
        "reconciliation_rows": [_stringify(row) for row in reconciliation],
        "signal_rows": signals,
        "band_rows": bands,
        "claim_ids": [row.get("Claim ID") for row in claim_table],
        "limitations": list(LIMITATIONS),
        "publication_label": INTERNAL_STATUS if INTERNAL_STATUS in joined else None,
        "heading_count": len(headings),
        "table_count": len(document.tables),
    }
    return receipt


def _validate_cross_artifact(
    excel: Mapping[str, Any],
    powerpoint: Mapping[str, Any],
    word: Mapping[str, Any],
    canonical: Mapping[str, Any],
    issues: list[OfficeIssue],
) -> None:
    for field in ("run_id", "definition_version", "source_sha256", "publication_label"):
        values = {artifact.get(field) for artifact in (excel, powerpoint, word)}
        expected = INTERNAL_STATUS if field == "publication_label" else canonical[field]
        if values != {expected}:
            _issue(
                issues,
                "cross_artifact_identity_mismatch",
                f"$.artifacts.{field}",
                f"Excel, PowerPoint, and Word do not share canonical {field}",
            )
    if excel.get("claim_ids") != list(REQUIRED_CLAIM_IDS) or word.get(
        "claim_ids"
    ) != list(REQUIRED_CLAIM_IDS):
        _issue(
            issues,
            "cross_artifact_claim_mismatch",
            "$.artifacts.claim_ids",
            "Excel and Word must carry all six claims",
        )


def _checks(
    packet: Mapping[str, Any],
    gate: int,
    issues: list[OfficeIssue],
) -> dict[str, Mapping[str, Any]]:
    value = packet.get("checks")
    if not isinstance(value, list):
        _issue(
            issues,
            "gate_checks_missing",
            f"$.validation.gate_packets[{gate}].checks",
            "checks must be a list",
        )
        return {}
    result: dict[str, Mapping[str, Any]] = {}
    for row in value:
        if not isinstance(row, Mapping) or not isinstance(row.get("check_id"), str):
            continue
        if row["check_id"] in result or row.get("status") != "pass":
            _issue(
                issues,
                "gate_check_invalid",
                f"$.validation.gate_packets[{gate}].checks",
                "gate checks must be unique and passing",
            )
        result[str(row["check_id"])] = row
    return result


def _evidence(
    checks: Mapping[str, Mapping[str, Any]],
    check_id: str,
    issues: list[OfficeIssue],
) -> Mapping[str, Any]:
    row = checks.get(check_id)
    evidence = row.get("evidence") if isinstance(row, Mapping) else None
    if not isinstance(evidence, Mapping):
        _issue(
            issues,
            "gate_evidence_missing",
            f"$.validation.checks.{check_id}",
            "required evidence object is missing",
        )
        return {}
    return evidence


def _reconciliation_rows(canonical: Mapping[str, Any]) -> list[dict[str, Any]]:
    return [
        {
            "Metric": "Source rows → raw rows",
            "Source / expected": canonical["source_rows"],
            "Derived / actual": canonical["raw_rows"],
            "Difference": 0,
            "Status": "PASS",
            "Claim ID": "W1-CLM-001",
            "Evidence": "Gate 1 raw_rows_reconciled",
        },
        {
            "Metric": "Raw rows → borrowers",
            "Source / expected": canonical["raw_rows"],
            "Derived / actual": canonical["borrower_rows"],
            "Difference": 0,
            "Status": "PASS",
            "Claim ID": "W1-CLM-002",
            "Evidence": "Gate 2 wide_to_long_rows",
        },
        {
            "Metric": "Borrowers × 6 → borrower-months",
            "Source / expected": canonical["borrower_rows"]
            * canonical["months_per_borrower"],
            "Derived / actual": canonical["account_month_rows"],
            "Difference": 0,
            "Status": "PASS",
            "Claim ID": "W1-CLM-002",
            "Evidence": "Gate 2 wide_to_long_rows",
        },
        {
            "Metric": "Wide bills → long bills",
            "Source / expected": canonical["bill_source_total"],
            "Derived / actual": canonical["bill_derived_total"],
            "Difference": canonical["bill_difference"],
            "Status": "PASS",
            "Claim ID": "W1-CLM-003",
            "Evidence": "Gate 2 bill_amounts",
        },
        {
            "Metric": "Wide payments → long payments",
            "Source / expected": canonical["payment_source_total"],
            "Derived / actual": canonical["payment_derived_total"],
            "Difference": canonical["payment_difference"],
            "Status": "PASS",
            "Claim ID": "W1-CLM-003",
            "Evidence": "Gate 2 payment_amounts",
        },
    ]


def _dq_rows(canonical: Mapping[str, Any]) -> list[dict[str, Any]]:
    checks = {
        gate: {
            row["check_id"]: row["evidence"]
            for row in packet["checks"]
        }
        for gate, packet in canonical["gate_packets"].items()
    }
    g2, g3 = checks[2], checks[3]
    return [
        {
            "Control ID": "DQ-PK-BORROWER",
            "Severity": "Critical",
            "Status": "PASS",
            "Observed": g2["primary_key"]["duplicate_count"],
            "Threshold / expectation": "0",
            "Affected output": "core.borrower",
            "Evidence / warning": "Gate 2 primary_key",
        },
        {
            "Control ID": "DQ-PK-ACCOUNT-MONTH",
            "Severity": "Critical",
            "Status": "PASS",
            "Observed": g2["primary_key"]["account_month_duplicate_count"],
            "Threshold / expectation": "0",
            "Affected output": "core.account_month",
            "Evidence / warning": "Gate 2 primary_key",
        },
        {
            "Control ID": "DQ-NULL-BORROWER-KEY",
            "Severity": "Critical",
            "Status": "PASS",
            "Observed": g2["primary_key"]["null_key_count"],
            "Threshold / expectation": "0",
            "Affected output": "core.borrower",
            "Evidence / warning": "Gate 2 primary_key",
        },
        {
            "Control ID": "DQ-FK-ACCOUNT-MONTH",
            "Severity": "Critical",
            "Status": "PASS",
            "Observed": g2["foreign_key"]["orphan_count"],
            "Threshold / expectation": "0",
            "Affected output": "core.account_month",
            "Evidence / warning": "Gate 2 foreign_key",
        },
        {
            "Control ID": "DQ-CODE-DOMAIN",
            "Severity": "High",
            "Status": "PASS",
            "Observed": g2["code_domains"]["invalid_count"],
            "Threshold / expectation": "0",
            "Affected output": "core tables",
            "Evidence / warning": "Gate 2 code_domains",
        },
        {
            "Control ID": "DQ-MISSING-HISTORY",
            "Severity": "High",
            "Status": "PASS",
            "Observed": g2["code_domains"]["missing_required_count"],
            "Threshold / expectation": "0",
            "Affected output": "core.account_month",
            "Evidence / warning": "Gate 2 code_domains",
        },
        {
            "Control ID": "DQ-SQL-FAILURES",
            "Severity": "Critical",
            "Status": "PASS",
            "Observed": g2["sql_test_suite"]["failed_count"],
            "Threshold / expectation": "0",
            "Affected output": "all Wave 1 marts",
            "Evidence / warning": "Gate 2 sql_test_suite",
        },
        {
            "Control ID": "DQ-TIMING",
            "Severity": "Critical",
            "Status": "PASS",
            "Observed": g3["observation_timing"]["violation_count"],
            "Threshold / expectation": "0",
            "Affected output": "risk signals",
            "Evidence / warning": "Gate 3 observation_timing",
        },
        {
            "Control ID": "DQ-LEAKAGE",
            "Severity": "Critical",
            "Status": "PASS",
            "Observed": g3["leakage"]["violation_count"],
            "Threshold / expectation": "0",
            "Affected output": "risk signals",
            "Evidence / warning": "Gate 3 leakage",
        },
        {
            "Control ID": "DQ-SQL-DETAIL-COVERAGE",
            "Severity": "Info",
            "Status": "LIMITATION",
            "Observed": f"{canonical['sql_test_count']} aggregate tests",
            "Threshold / expectation": "individual test IDs remain upstream",
            "Affected output": "test-level audit drilldown",
            "Evidence / warning": (
                "validation input exposes counts but not individual test records"
            ),
        },
    ]


def _excel_signal_rows(canonical: Mapping[str, Any]) -> list[dict[str, Any]]:
    return [
        {
            "Signal ID": row["signal_id"],
            "Bucket": row["bucket"],
            "Sample count": row["sample_count"],
            "Default count": row["default_count"],
            "Observed default rate exact": row["observed_default_rate"],
            "Observed default rate": float(row["observed_default_rate"]),
            "Claim ID": "W1-CLM-005",
        }
        for row in canonical["signal_rows"]
    ]


def _excel_band_rows(canonical: Mapping[str, Any]) -> list[dict[str, Any]]:
    return [
        {
            "Risk band": row["risk_band"],
            "Sample count": row["sample_count"],
            "Default count": row["default_count"],
            "Observed default rate exact": row["observed_default_rate"],
            "Observed default rate": float(row["observed_default_rate"]),
            "Claim ID": "W1-CLM-006",
        }
        for row in canonical["band_rows"]
    ]


def _excel_gate_rows(canonical: Mapping[str, Any]) -> list[dict[str, Any]]:
    result = []
    for gate in (1, 2, 3):
        for row in canonical["gate_packets"][gate]["checks"]:
            result.append(
                {
                    "Gate": gate,
                    "Check ID": row["check_id"],
                    "Status": row["status"],
                    "Evidence JSON": json.dumps(
                        row["evidence"], ensure_ascii=False, sort_keys=True
                    ),
                }
            )
    return result


def _excel_claim_rows(
    claims: Sequence[Mapping[str, Any]],
    canonical: Mapping[str, Any],
    validation_path: Path,
) -> list[dict[str, Any]]:
    return [
        {
            "Claim ID": row["claim_id"],
            "Claim": row["claim"],
            "Run ID": canonical["run_id"],
            "Validation SHA-256": _sha256(validation_path),
            "JSON pointers": " | ".join(row["json_pointers"]),
        }
        for row in claims
    ]


def _sheet_dicts(sheet: Any) -> list[dict[str, Any]]:
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(value) if value is not None else "" for value in rows[0]]
    return [
        dict(zip(headers, row, strict=True))
        for row in rows[1:]
        if any(value is not None for value in row)
    ]


def _doc_table_dicts(table: Any) -> list[dict[str, str]]:
    headers = [cell.text for cell in table.rows[0].cells]
    return [
        dict(zip(headers, [cell.text for cell in row.cells], strict=True))
        for row in table.rows[1:]
    ]


def _ppt_table(tables: Sequence[list[list[str]]], first_header: str) -> list[list[str]]:
    for table in tables:
        if table and table[0] and table[0][0] == first_header:
            return table
    return []


def _gate_5_packet(
    run_id: str,
    generated_at: str,
    *,
    office_manifest_path: Path,
    evidence_map_path: Path,
    claim_count: int,
    artifact_count: int,
    limitation_count: int,
) -> dict[str, Any]:
    return {
        "schema_version": "1.0",
        "gate": 5,
        "run_id": run_id,
        "generated_at": generated_at,
        "producer": {
            "role": "Validation",
            "command": (
                "python -m src.validation.wave1_office_validation "
                "--office-manifest <manifest> --validation <validation> "
                "--claim-manifest <claims> --evidence-map <map> "
                "--output <report> --gate-output <gate_5.json>"
            ),
        },
        "status": "pass",
        "checks": [
            {
                "check_id": "shared_run_id",
                "status": "pass",
                "evidence": {
                    "distinct_run_id_count": 1,
                    "run_id": run_id,
                    "artifact_count": artifact_count,
                },
            },
            {
                "check_id": "cross_artifact_values",
                "status": "pass",
                "evidence": {
                    "mismatch_count": 0,
                    "artifact_count": artifact_count,
                    "office_manifest": _portable(office_manifest_path),
                },
            },
            {
                "check_id": "limitations",
                "status": "pass",
                "evidence": {
                    "documented": True,
                    "limitation_count": limitation_count,
                    "human_publication_approval": "pending",
                },
            },
            {
                "check_id": "evidence_map",
                "status": "pass",
                "evidence": {
                    "unresolved_claim_count": 0,
                    "claim_count": claim_count,
                    "evidence_map": _portable(evidence_map_path),
                },
            },
        ],
    }


def _report(
    run_id: str | None,
    generated_at: str,
    issues: Sequence[OfficeIssue],
    unresolved_claim_count: int,
    limitations_documented: bool,
    artifact_checks: Mapping[str, Any],
    gate_packet: Mapping[str, Any] | None,
) -> OfficeValidationReport:
    libreoffice = shutil.which("libreoffice") or shutil.which("soffice")
    return OfficeValidationReport(
        schema_version="1.0",
        run_id=run_id,
        generated_at=generated_at,
        status=(
            "pass"
            if not issues and unresolved_claim_count == 0 and gate_packet is not None
            else "fail"
        ),
        mismatch_count=len(issues),
        unresolved_claim_count=unresolved_claim_count,
        limitations_documented=limitations_documented,
        human_publication_approval="pending",
        visual_render={
            "status": "not_run",
            "reason": (
                "LibreOffice executable unavailable; structural OOXML QA completed"
                if libreoffice is None
                else "structural OOXML QA completed; visual rendering is a separate review"
            ),
            "libreoffice_executable": libreoffice,
            "structural_qa": "completed",
        },
        artifact_checks=dict(artifact_checks),
        issues=tuple(issues),
        gate_packet=gate_packet,
    )


def _match(
    actual: object,
    expected: object,
    issues: list[OfficeIssue],
    code: str,
    path: str,
) -> None:
    if actual != expected:
        _issue(
            issues,
            code,
            path,
            "actual Office content does not exactly match canonical evidence",
        )


def _match_excel_outcomes(
    actual: object,
    expected: object,
    issues: list[OfficeIssue],
    code: str,
    path: str,
) -> None:
    """Compare exact audit fields and tolerate only OOXML binary-float encoding."""

    if not isinstance(actual, list) or not isinstance(expected, list) or len(actual) != len(
        expected
    ):
        _issue(issues, code, path, "Office outcome row coverage does not match")
        return
    for actual_row, expected_row in zip(actual, expected, strict=True):
        if not isinstance(actual_row, Mapping) or not isinstance(expected_row, Mapping):
            _issue(issues, code, path, "Office outcome row is not an object")
            return
        actual_exact = dict(actual_row)
        expected_exact = dict(expected_row)
        actual_display = actual_exact.pop("Observed default rate", None)
        expected_display = expected_exact.pop("Observed default rate", None)
        if actual_exact != expected_exact:
            _issue(
                issues,
                code,
                path,
                "exact count/rate-string fields do not match canonical evidence",
            )
            return
        try:
            within_binary_encoding = (
                abs(float(actual_display) - float(expected_display)) <= 1e-15
            )
        except (TypeError, ValueError):
            within_binary_encoding = False
        if not within_binary_encoding:
            _issue(
                issues,
                code,
                path,
                "display-rate float differs beyond fixed OOXML encoding tolerance",
            )
            return


def _stringify(row: Mapping[str, Any]) -> dict[str, str]:
    return {key: str(value) for key, value in row.items()}


def _integer(value: object) -> str:
    return f"{int(value):,}"


def _percent(value: object) -> str:
    from decimal import Decimal, ROUND_HALF_UP

    return f"{(Decimal(str(value)) * Decimal('100')).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)}%"


def _resolve_portable(
    value: object,
    root: Path,
    issues: list[OfficeIssue],
) -> Path | None:
    if not isinstance(value, str) or not value:
        _issue(issues, "invalid_path", "$.path", "path must be a repository-relative string")
        return None
    relative = Path(value)
    if relative.is_absolute():
        _issue(issues, "absolute_path", value, "manifest paths must be repository-relative")
        return None
    resolved = (root / relative).resolve()
    try:
        resolved.relative_to(root)
    except ValueError:
        _issue(issues, "path_escape", value, "manifest path escapes repository root")
        return None
    return resolved


def _portable(path: Path) -> str:
    try:
        return str(path.resolve().relative_to(Path.cwd().resolve()))
    except ValueError:
        return str(path)


def _portable_to_root(path: Path, root: Path) -> str:
    try:
        return str(path.resolve().relative_to(root.resolve()))
    except ValueError:
        return str(path)


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _issue(
    issues: list[OfficeIssue],
    code: str,
    path: str,
    message: str,
) -> None:
    issues.append(OfficeIssue(code, path, message))


def _write_new_json(path: Path, value: Mapping[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("x", encoding="utf-8") as handle:
        json.dump(value, handle, indent=2, ensure_ascii=False, allow_nan=False)
        handle.write("\n")


def _parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--office-manifest", type=Path, required=True)
    parser.add_argument("--validation", type=Path, required=True)
    parser.add_argument("--claim-manifest", type=Path, required=True)
    parser.add_argument("--evidence-map", type=Path, required=True)
    parser.add_argument("--artifact-root", type=Path, default=Path.cwd())
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--gate-output", type=Path)
    return parser


def main(argv: Sequence[str] | None = None) -> int:
    args = _parser().parse_args(argv)
    if args.output.exists() or (
        args.gate_output is not None and args.gate_output.exists()
    ):
        print("Office validation blocked: refusing to overwrite validation evidence")
        return 2
    try:
        report = validate_office_pack(
            args.office_manifest,
            args.validation,
            args.claim_manifest,
            args.evidence_map,
            artifact_root=args.artifact_root,
        )
        _write_new_json(args.output, report.to_dict())
        if report.status == "pass" and args.gate_output is not None:
            if report.gate_packet is None:
                raise ValueError("passing report is missing Gate 5 packet")
            _write_new_json(args.gate_output, report.gate_packet)
    except (OSError, ValueError, json.JSONDecodeError) as exc:
        print(f"Office validation blocked: {exc}")
        return 2
    print(
        f"Office validation {report.status}: run_id={report.run_id!r}, "
        f"mismatches={report.mismatch_count}, "
        f"unresolved_claims={report.unresolved_claim_count}"
    )
    return 0 if report.status == "pass" else 1


if __name__ == "__main__":
    raise SystemExit(main())
